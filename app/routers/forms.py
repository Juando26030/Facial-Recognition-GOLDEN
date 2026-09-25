"""Formularios Web — lado del coordinador (Sprint 5): crear/editar formularios de un evento, estados y calendario, cupo,
plantillas, respuestas, archivos, invitaciones, reporte y carga a la base. Todo coordinador+ (comercial incluida).
La página pública que llena la gente vive en `forms_public.py` (`/f/<evento>/<formulario>`, sin login)."""
import json
import os
import secrets
import tempfile
from datetime import datetime
from typing import Optional

from fastapi import APIRouter, Depends, File, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app import bulk_jobs, formlib, formsvc, wompi
from app.auth import get_event_for_staff, require_role
from app.database import get_db
from app.mailer import send_mail
from app.models import Event, FormEvent, FormInvite, FormPayment, FormPerson, FormRefund, FormSubmission, SavedFormTemplate, StaffUser, WebForm
from app.routers import parametros
from app.routers.badges import ALLOWED_IMAGE_EXT, _badge_assets_dir
from app.timeutil import to_local

router = APIRouter()
STAFF = require_role("coordinador")


# ------------------------------------------------------------------ utilidades
def _get_form(db: Session, event: Event, form_id: int) -> WebForm:
    form = db.query(WebForm).filter(WebForm.id == form_id, WebForm.event_id == event.id).first()
    if not form:
        raise HTTPException(status_code=404, detail="Formulario no encontrado")
    return form


def _base_url(request: Request) -> str:
    return (os.getenv("PUBLIC_BASE_URL") or str(request.base_url)).rstrip("/")


def _urls(request: Request, form: WebForm) -> dict:
    base = f"{_base_url(request)}/f/{form.event_id}/{form.slug}"
    return {"public_url": base, "test_url": f"{base}?k={form.test_key}"}


def _unique_slug(db: Session, event_id: int, wanted: str, exclude_id: Optional[int] = None) -> str:
    base = formlib.slugify(wanted)
    slug, n = base, 2
    while True:
        q = db.query(WebForm).filter(WebForm.event_id == event_id, WebForm.slug == slug)
        if exclude_id:
            q = q.filter(WebForm.id != exclude_id)
        if not q.first():
            return slug
        slug, n = f"{base}-{n}", n + 1


def _optional_keys(event: Event) -> set:
    return set(event.get_optional_labels().keys())


def _allocate_event_fields(event: Event, design: dict) -> dict:
    """Campos marcados «guardar también en Parámetros del Evento» (`sync`) y sin vínculo: se les asigna el primer
    `opcional_N` libre y se rotula con su nombre — mismo mecanismo que «agregar campo» en Parámetros."""
    labels = event.get_optional_labels()
    for f in design["fields"].values():
        if f["type"] in formlib.INPUT_TYPES and f.get("sync") and not f.get("key"):
            used = {int(k.split("_")[1]) for k in labels}
            slot = next((n for n in range(1, 31) if n not in used), None)
            if slot is None:
                raise HTTPException(status_code=400, detail="Ya se usaron los 30 campos opcionales del evento")
            key = f"opcional_{slot}"
            labels[key] = f["label"]
            f["key"], f["sync"] = key, False
    event.set_optional_labels(labels)
    return design


def _summary(db: Session, form: WebForm, request: Request) -> dict:
    real = formsvc.real_submissions(db, form).count()
    tests = db.query(FormSubmission).filter(FormSubmission.form_id == form.id, FormSubmission.is_test == True, FormSubmission.status == "confirmed").count()  # noqa: E712
    return {
        "id": form.id, "name": form.name, "slug": form.slug, "manual_status": form.manual_status, "use_schedule": form.use_schedule,
        "status": formsvc.status_of(form), "public_state": formsvc.public_state(db, form), "capacity": form.capacity,
        "submissions": real, "test_submissions": tests, "created_at": to_local(form.created_at).strftime("%Y-%m-%d %H:%M"),
        "feed": formsvc.get_settings(form)["feed"], **_urls(request, form),
    }


def _detail(db: Session, form: WebForm, request: Request) -> dict:
    return {**_summary(db, form, request), "design": formsvc.get_design(form), "settings": formsvc.get_settings(form),
            "schedule": formsvc.get_schedule(form), "test_key": form.test_key, "fed_at": form.fed_at.isoformat() if form.fed_at else None,
            "payments": {"has_field": bool(formlib.payment_field(formsvc.get_design(form))), "sandbox_configured": bool(wompi.config(True)), "production_configured": bool((wompi.config(False) or {}).get("test") is False)}}


# ------------------------------------------------------------------ CRUD
@router.get("/events/{event_id}/forms")
async def list_forms(event_id: int, request: Request, db: Session = Depends(get_db), staff: StaffUser = Depends(STAFF)):
    event = get_event_for_staff(event_id, db, staff)
    forms = db.query(WebForm).filter(WebForm.event_id == event.id).order_by(WebForm.id.desc()).all()
    for f in forms:
        _maybe_feed_on_close(db, f)
    return [_summary(db, f, request) for f in forms]


@router.post("/events/{event_id}/forms")
async def create_form(event_id: int, data: dict, request: Request, db: Session = Depends(get_db), staff: StaffUser = Depends(STAFF)):
    """Crea un formulario nuevo (con campos de partida) o, con `template_id`, a partir de una plantilla guardada."""
    event = get_event_for_staff(event_id, db, staff)
    name = str(data.get("name") or "").strip()[:120]
    if not name:
        raise HTTPException(status_code=400, detail="Ponle un nombre al formulario")
    if data.get("template_id"):
        tpl = db.query(SavedFormTemplate).filter_by(id=data["template_id"], tenant_id=event.tenant_id).first()
        if not tpl:
            raise HTTPException(status_code=404, detail="Plantilla no encontrada")
        design = json.loads(tpl.design_json)
        design["theme"]["title"] = design["theme"].get("title") or name
    else:
        design = formlib.default_design(name)
    try:
        design = formlib.sanitize_design(design, _optional_keys(event) | {f"opcional_{i}" for i in range(1, 31)})
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    # un diseño importado puede traer campos opcional_N que este evento no tiene: se desvinculan (quedan como campos solo del formulario)
    valid = set(formlib.IDENTITY_KEYS) | _optional_keys(event)
    for f in design["fields"].values():
        if f.get("key") and f["key"] not in valid:
            f["key"] = None
    form = WebForm(
        event_id=event.id, name=name, slug=_unique_slug(db, event.id, name), manual_status="pruebas", design_json=json.dumps(design),
        settings_json=json.dumps(formlib.sanitize_settings({})), test_key=secrets.token_urlsafe(12), created_by_id=staff.id,
    )
    db.add(form)
    db.commit()
    return _detail(db, form, request)


@router.get("/events/{event_id}/forms/{form_id}")
async def get_form(event_id: int, form_id: int, request: Request, db: Session = Depends(get_db), staff: StaffUser = Depends(STAFF)):
    event = get_event_for_staff(event_id, db, staff)
    return _detail(db, _get_form(db, event, form_id), request)


@router.put("/events/{event_id}/forms/{form_id}")
async def update_form(event_id: int, form_id: int, data: dict, request: Request, db: Session = Depends(get_db), staff: StaffUser = Depends(STAFF)):
    """Guarda nombre, dirección (slug), diseño y/o configuración. El diseño se valida completo en el servidor."""
    event = get_event_for_staff(event_id, db, staff)
    form = _get_form(db, event, form_id)
    if "name" in data:
        name = str(data["name"]).strip()[:120]
        if not name:
            raise HTTPException(status_code=400, detail="El formulario necesita un nombre")
        form.name = name
    if "slug" in data:
        slug = formlib.slugify(data["slug"])
        if db.query(WebForm).filter(WebForm.event_id == event.id, WebForm.slug == slug, WebForm.id != form.id).first():
            raise HTTPException(status_code=400, detail="Ya hay otro formulario con esa dirección en este evento")
        form.slug = slug
    try:
        if "design" in data:
            design = formlib.sanitize_design(data["design"], _optional_keys(event))
            design = _allocate_event_fields(event, design)
            form.design_json = json.dumps(design)
        if "settings" in data:
            form.settings_json = json.dumps(formlib.sanitize_settings(data["settings"]))
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    db.commit()
    return _detail(db, form, request)


@router.delete("/events/{event_id}/forms/{form_id}")
async def delete_form(event_id: int, form_id: int, db: Session = Depends(get_db), staff: StaffUser = Depends(STAFF)):
    event = get_event_for_staff(event_id, db, staff)
    form = _get_form(db, event, form_id)
    delete_form_rows(db, form)
    db.commit()
    return {"message": "Formulario eliminado con sus inscripciones"}


def delete_form_rows(db: Session, form: WebForm) -> None:
    import shutil
    if db.query(FormPayment).filter(FormPayment.form_id == form.id, FormPayment.status.in_(("approved", "refunded")), FormPayment.is_test == False).first():  # noqa: E712
        raise HTTPException(status_code=409, detail=f"El formulario «{form.name}» tiene pagos aprobados: son registros financieros y no se pueden borrar. Ciérralo o finalízalo en su lugar.")
    db.query(FormRefund).filter(FormRefund.payment_id.in_(db.query(FormPayment.id).filter(FormPayment.form_id == form.id))).delete(synchronize_session=False)
    db.query(FormPayment).filter(FormPayment.form_id == form.id).delete()
    db.query(FormSubmission).filter(FormSubmission.form_id == form.id).delete()
    db.query(FormEvent).filter(FormEvent.form_id == form.id).delete()
    db.query(FormInvite).filter(FormInvite.form_id == form.id).delete()
    db.query(FormPerson).filter(FormPerson.form_id == form.id).delete()
    event = db.query(Event).filter(Event.id == form.event_id).first()
    if event:
        shutil.rmtree(os.path.join("data", event.tenant_id, "form_files", str(form.id)), ignore_errors=True)
    db.delete(form)


@router.post("/events/{event_id}/forms/{form_id}/duplicate")
async def duplicate_form(event_id: int, form_id: int, request: Request, db: Session = Depends(get_db), staff: StaffUser = Depends(STAFF)):
    event = get_event_for_staff(event_id, db, staff)
    src = _get_form(db, event, form_id)
    name = f"{src.name} (copia)"
    copy = WebForm(event_id=event.id, name=name, slug=_unique_slug(db, event.id, name), manual_status="pruebas", design_json=src.design_json,
                   settings_json=src.settings_json, capacity=src.capacity, test_key=secrets.token_urlsafe(12), created_by_id=staff.id)
    db.add(copy)
    db.commit()
    return _detail(db, copy, request)


# ------------------------------------------------------------------ estado, calendario y cupo
@router.put("/events/{event_id}/forms/{form_id}/status")
async def set_status(event_id: int, form_id: int, data: dict, request: Request, db: Session = Depends(get_db), staff: StaffUser = Depends(STAFF)):
    """Estado manual, calendario por fechas y cupo (editable en caliente: no se pierde ninguna inscripción)."""
    event = get_event_for_staff(event_id, db, staff)
    form = _get_form(db, event, form_id)
    if "manual_status" in data:
        if data["manual_status"] not in formlib.STATUSES:
            raise HTTPException(status_code=400, detail=f"Estado inválido (uno de: {', '.join(formlib.STATUSES)})")
        form.manual_status = data["manual_status"]
    if "schedule" in data:
        try:
            form.schedule_json = json.dumps(formlib.parse_schedule(data["schedule"]))
        except ValueError as e:
            raise HTTPException(status_code=400, detail=str(e))
    if "use_schedule" in data:
        form.use_schedule = bool(data["use_schedule"])
        if form.use_schedule and not formsvc.get_schedule(form):
            raise HTTPException(status_code=400, detail="Para usar el calendario define al menos un tramo")
    if "capacity" in data:
        cap = data["capacity"]
        if cap in (None, ""):
            form.capacity = None
        else:
            try:
                cap = int(cap)
            except (TypeError, ValueError):
                raise HTTPException(status_code=400, detail="El cupo debe ser un número")
            if cap < 1:
                raise HTTPException(status_code=400, detail="El cupo debe ser al menos 1")
            form.capacity = cap
    db.commit()
    _maybe_feed_on_close(db, form)
    return _detail(db, form, request)


def _upsert():
    from app.routers.api import _upsert_attendee
    return _upsert_attendee


def _maybe_feed_on_close(db: Session, form: WebForm) -> None:
    """Modo «al cerrar»: cuando el formulario ya está cerrado/finalizado, se cargan solas todas las inscripciones
    pendientes (una vez; las que lleguen después —no debería— se cargan con «Cargar ahora»)."""
    if formsvc.get_settings(form)["feed"] == "on_close" and form.fed_at is None and formsvc.status_of(form) in ("cerrado", "finalizado"):
        formsvc.feed_pending(db, form, _upsert())


@router.post("/events/{event_id}/forms/{form_id}/feed-now")
async def feed_now(event_id: int, form_id: int, db: Session = Depends(get_db), staff: StaffUser = Depends(STAFF)):
    """Carga a la base del evento todas las inscripciones reales aún no cargadas (quedan «No registrado»)."""
    event = get_event_for_staff(event_id, db, staff)
    form = _get_form(db, event, form_id)
    return formsvc.feed_pending(db, form, _upsert())


# ------------------------------------------------------------------ campos del evento (Parámetros) y plantillas
@router.get("/events/{event_id}/form-event-fields")
async def event_fields(event_id: int, db: Session = Depends(get_db), staff: StaffUser = Depends(STAFF)):
    """Campos del evento (Parámetros) disponibles para arrastrar al formulario: los de identidad y cada opcional."""
    event = get_event_for_staff(event_id, db, staff)
    out = []
    for cfg in parametros.field_configs_for_event(db, event):
        key = cfg["key"]
        if key in ("categories", "certificate", "digital_contact") or cfg["field_type"] == "signature":
            continue
        kind = {"boolean": "checkbox", "consent": "checkbox", "email": "email"}.get(cfg["field_type"], cfg["field_type"])
        if key == "email":
            kind = "email"
        if key == "phone":
            kind = "phone"
        out.append({"key": key, "label": cfg["label"], "type": kind if kind in formlib.INPUT_TYPES else "text_short", "options": cfg.get("options") or [], "required": bool(cfg.get("required"))})
    return out


@router.get("/events/{event_id}/form-templates")
async def list_templates(event_id: int, db: Session = Depends(get_db), staff: StaffUser = Depends(STAFF)):
    event = get_event_for_staff(event_id, db, staff)
    rows = db.query(SavedFormTemplate).filter_by(tenant_id=event.tenant_id).order_by(SavedFormTemplate.name).all()
    return [{"id": t.id, "name": t.name, "fields": len(json.loads(t.design_json)["fields"])} for t in rows]


@router.post("/events/{event_id}/forms/{form_id}/save-as-template")
async def save_template(event_id: int, form_id: int, data: dict, db: Session = Depends(get_db), staff: StaffUser = Depends(STAFF)):
    event = get_event_for_staff(event_id, db, staff)
    form = _get_form(db, event, form_id)
    name = str(data.get("name") or "").strip()[:120]
    if not name:
        raise HTTPException(status_code=400, detail="Ponle un nombre a la plantilla")
    tpl = db.query(SavedFormTemplate).filter_by(tenant_id=event.tenant_id, name=name).first()
    if tpl:
        tpl.design_json = form.design_json          # mismo nombre: se actualiza
    else:
        tpl = SavedFormTemplate(tenant_id=event.tenant_id, name=name, design_json=form.design_json)
        db.add(tpl)
    db.commit()
    return {"id": tpl.id, "name": tpl.name}


@router.delete("/events/{event_id}/form-templates/{template_id}")
async def delete_template(event_id: int, template_id: int, db: Session = Depends(get_db), staff: StaffUser = Depends(STAFF)):
    event = get_event_for_staff(event_id, db, staff)
    tpl = db.query(SavedFormTemplate).filter_by(id=template_id, tenant_id=event.tenant_id).first()
    if not tpl:
        raise HTTPException(status_code=404, detail="Plantilla no encontrada")
    db.delete(tpl)
    db.commit()
    return {"message": "Plantilla eliminada de la librería"}


@router.post("/events/{event_id}/forms-upload-image")
async def upload_image(event_id: int, file: UploadFile = File(...), db: Session = Depends(get_db), staff: StaffUser = Depends(STAFF)):
    event = get_event_for_staff(event_id, db, staff)
    ext = os.path.splitext(file.filename or "")[1].lower()
    if ext not in ALLOWED_IMAGE_EXT:
        raise HTTPException(status_code=400, detail="Usa una imagen PNG, JPG, WEBP o GIF")
    content = await file.read()
    if len(content) > 8_000_000:
        raise HTTPException(status_code=400, detail="La imagen pesa más de 8 MB")
    name = f"{secrets.token_hex(16)}{ext}"
    with open(os.path.join(_badge_assets_dir(event.tenant_id), name), "wb") as f:
        f.write(content)
    return {"storage_path": f"{event.tenant_id}/{name}"}


# ------------------------------------------------------------------ respuestas y archivos
def _column_defs(design: dict) -> list:
    """(id, etiqueta) de los campos que guardan datos, en el orden en que aparecen en el formulario."""
    out = []
    for row in design["rows"]:
        for fid in row["items"]:
            f = design["fields"][fid]
            if f["type"] in formlib.INPUT_TYPES:
                out.append((fid, f["label"], f["type"]))
    return out


def _display(value):
    if isinstance(value, dict):
        return value.get("filename", "")
    if isinstance(value, list):
        return ", ".join(value)
    if value is True:
        return "Sí"
    if value is False:
        return "No"
    return "" if value is None else value


@router.get("/events/{event_id}/forms/{form_id}/submissions")
async def list_submissions(event_id: int, form_id: int, include_tests: bool = True, db: Session = Depends(get_db), staff: StaffUser = Depends(STAFF)):
    event = get_event_for_staff(event_id, db, staff)
    form = _get_form(db, event, form_id)
    design = formsvc.get_design(form)
    q = db.query(FormSubmission).filter(FormSubmission.form_id == form.id, FormSubmission.status == "confirmed")
    if not include_tests:
        q = q.filter(FormSubmission.is_test == False)  # noqa: E712
    rows = q.order_by(FormSubmission.id.desc()).limit(2000).all()
    paid = _approved_by_submission(db, form)
    return {"has_payment": bool(formlib.payment_field(design)), "columns": [{"id": i, "label": l, "type": t} for i, l, t in _column_defs(design)], "rows": [
        {"id": s.id, "created_at": to_local(s.created_at).strftime("%Y-%m-%d %H:%M:%S"), "is_test": s.is_test, "fed": s.fed, "source": s.source or "",
         "paid": paid[s.id].amount_cents // 100 if s.id in paid else None, "data": {k: _display(v) for k, v in json.loads(s.data_json).items()}} for s in rows]}


@router.delete("/events/{event_id}/forms/{form_id}/submissions/{submission_id}")
async def delete_submission(event_id: int, form_id: int, submission_id: int, db: Session = Depends(get_db), staff: StaffUser = Depends(STAFF)):
    event = get_event_for_staff(event_id, db, staff)
    form = _get_form(db, event, form_id)
    sub = db.query(FormSubmission).filter_by(id=submission_id, form_id=form.id).first()
    if not sub:
        raise HTTPException(status_code=404, detail="Inscripción no encontrada")
    for v in json.loads(sub.data_json).values():
        if isinstance(v, dict) and v.get("stored"):
            try:
                os.remove(os.path.join(formsvc.form_files_dir(event.tenant_id, form.id), v["stored"]))
            except OSError:
                pass
    db.query(FormPayment).filter(FormPayment.submission_id == sub.id).update({"submission_id": None}, synchronize_session=False)
    db.delete(sub)
    db.commit()
    return {"message": "Inscripción eliminada (si tenía un pago, ese registro financiero se conserva)"}


@router.get("/events/{event_id}/forms/{form_id}/files/{submission_id}/{field_id}")
async def download_file(event_id: int, form_id: int, submission_id: int, field_id: str, db: Session = Depends(get_db), staff: StaffUser = Depends(STAFF)):
    event = get_event_for_staff(event_id, db, staff)
    form = _get_form(db, event, form_id)
    sub = db.query(FormSubmission).filter_by(id=submission_id, form_id=form.id).first()
    info = json.loads(sub.data_json).get(field_id) if sub else None
    if not isinstance(info, dict) or not info.get("stored"):
        raise HTTPException(status_code=404, detail="Archivo no encontrado")
    path = os.path.join(formsvc.form_files_dir(event.tenant_id, form.id), os.path.basename(info["stored"]))
    if not os.path.isfile(path):
        raise HTTPException(status_code=404, detail="Archivo no encontrado")
    return FileResponse(path, filename=info.get("filename") or os.path.basename(path))


@router.get("/events/{event_id}/forms/{form_id}/report")
async def report(event_id: int, form_id: int, include_tests: bool = False, request: Request = None, db: Session = Depends(get_db), staff: StaffUser = Depends(STAFF)):
    """Reporte Excel del formulario (mismas convenciones que `app/reports.py`: título arriba, encabezado oscuro,
    autofiltro). Por defecto SIN las inscripciones de prueba."""
    event = get_event_for_staff(event_id, db, staff)
    form = _get_form(db, event, form_id)
    design = formsvc.get_design(form)
    cols = _column_defs(design)
    q = db.query(FormSubmission).filter(FormSubmission.form_id == form.id, FormSubmission.status == "confirmed")
    if not include_tests:
        q = q.filter(FormSubmission.is_test == False)  # noqa: E712
    subs = q.order_by(FormSubmission.id).all()
    has_pay = bool(formlib.payment_field(design))
    paid = _approved_by_submission(db, form) if has_pay else {}
    wb = Workbook()
    ws = wb.active
    ws.title = "Inscripciones"
    ws.append([f"{form.name} — {event.name} ({event.event_code})"])
    ws["A1"].font = Font(bold=True, size=13, color="0A0E2E")
    headers = ["N°", "Fecha", "Hora"] + [l for _, l, _ in cols] + ["Origen", "Prueba", "En la base del evento"] + (["Monto pagado (COP)", "Referencia de pago", "Método de pago", "ID transacción Wompi", "Reglas y descuentos", "Reembolsado (COP)", "Comisión Wompi estimada (COP)", "Neto estimado de comisión (COP)"] if has_pay else [])
    ws.append(headers)
    for c in ws[2]:
        c.font = Font(color="FFFFFF", bold=True)
        c.fill = PatternFill(start_color="0A0E2E", end_color="0A0E2E", fill_type="solid")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for n, s in enumerate(subs, start=1):
        data = json.loads(s.data_json)
        local = to_local(s.created_at)
        ws.append([n, local.strftime("%Y-%m-%d"), local.strftime("%H:%M:%S")] + [_display(data.get(fid)) for fid, _, _ in cols] + [s.source or "Directo", "Sí" if s.is_test else "No", "Sí" if s.fed else "No"] + (_pay_cells(paid.get(s.id)) if has_pay else []))
    ws.auto_filter.ref = f"A2:{get_column_letter(len(headers))}{max(ws.max_row, 2)}"
    for i in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 24 if i > 3 else 10
    ws.freeze_panes = "A3"
    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
    wb.save(tmp.name)
    return FileResponse(tmp.name, filename=f"Formulario_{form.slug}.xlsx")


def _approved_by_submission(db: Session, form: WebForm) -> dict:
    return {p.submission_id: p for p in db.query(FormPayment).filter(FormPayment.form_id == form.id, FormPayment.status == "approved", FormPayment.submission_id != None)}  # noqa: E711


def _pay_cells(p: Optional[FormPayment]) -> list:
    if not p:
        return ["", "", "", "", "", "", "", ""]
    applied = "; ".join(f"{a['label']} ({a['effect']})" for a in json.loads(p.breakdown_json or "{}").get("applied", []))
    return [p.amount_cents // 100, p.reference, p.payment_method or "", p.transaction_id or "", applied, (p.refunded_cents or 0) // 100,
            _fee_cop(p), _net_cop(p)]


def _fee_cop(p: FormPayment) -> int:
    """Comisión Wompi ESTIMADA de un pago, en pesos (ver wompi.estimate_fee)."""
    return wompi.estimate_fee(p.amount_cents, p.payment_method or "") // 100


def _net_cop(p: FormPayment) -> int:
    """Lo que queda de un pago aprobado: pagado − reembolsado − comisión estimada (la comisión no se devuelve al reembolsar)."""
    return p.amount_cents // 100 - (p.refunded_cents or 0) // 100 - _fee_cop(p)


def _auto_refund_possible(p: FormPayment) -> bool:
    """¿Hay reembolso automático por API para este pago? Tarjeta (anulación) siempre; otros medios solo con la API V2 (sandbox)."""
    if (p.payment_method or "").upper() == "CARD":
        return True
    cfg = wompi.config(p.is_test)
    return bool(cfg and wompi.refunds_v2_enabled(cfg))


def _auto_partial_possible(p: FormPayment) -> bool:
    """Reembolso parcial automático: solo con la API V2 y fuera de tarjeta (la anulación de tarjeta es solo por el valor completo)."""
    if (p.payment_method or "").upper() == "CARD":
        return False
    return _auto_refund_possible(p)


def _pay_label(p: FormPayment) -> str:
    if p.status == "pending" and p.created_at < datetime.utcnow() - formsvc.PENDING_HOLD:
        return "abandoned"      # pendiente y ya pasó el tiempo de espera: la persona no terminó de pagar
    return p.status


# ------------------------------------------------------------------ pagos (campo «Pago», Wompi)
@router.get("/events/{event_id}/forms/{form_id}/payments")
async def list_payments(event_id: int, form_id: int, include_tests: bool = False, db: Session = Depends(get_db), staff: StaffUser = Depends(STAFF)):
    """Pagos del formulario con su referencia (evento + formulario + inscripción): la conciliación por evento vive aquí, no
    en el dashboard de Wompi. `orphan` = aprobado pero sin inscripción (revisar a mano)."""
    event = get_event_for_staff(event_id, db, staff)
    form = _get_form(db, event, form_id)
    q = db.query(FormPayment).filter(FormPayment.form_id == form.id)
    if not include_tests:
        q = q.filter(FormPayment.is_test == False)  # noqa: E712
    rows = q.order_by(FormPayment.id.desc()).limit(2000).all()
    paid = [p for p in rows if p.status in ("approved", "refunded")]
    refunds = {}
    for r in db.query(FormRefund).filter(FormRefund.payment_id.in_([p.id for p in rows] or [0])).order_by(FormRefund.id):
        refunds.setdefault(r.payment_id, []).append(r)
    return {
        "total_cop": (sum(p.amount_cents for p in paid) - sum(p.refunded_cents or 0 for p in paid)) // 100, "approved": len(paid),
        "wompi_status": wompi.service_status(), "fees_cop": sum(_fee_cop(p) for p in paid),
        "net_after_fees_cop": sum(_net_cop(p) for p in paid),
        "refunded_cop": sum(p.refunded_cents or 0 for p in paid) // 100,
        "rows": [{"id": p.id, "reference": p.reference, "amount": p.amount_cents // 100, "status": _pay_label(p), "method": p.payment_method or "", "transaction_id": p.transaction_id or "",
                  "person_id": p.person_id or "", "is_test": p.is_test, "submission_id": p.submission_id, "orphan": p.status in ("approved", "refunded") and not p.submission_id and not (p.refunded_cents or 0),
                  "fee": _fee_cop(p) if p.status in ("approved", "refunded") else 0,
                  "net": _net_cop(p) if p.status in ("approved", "refunded") else 0,
                  "refunded": (p.refunded_cents or 0) // 100, "remaining": (p.amount_cents - (p.refunded_cents or 0)) // 100 if p.status == "approved" else 0,
                  "auto_refund": p.status == "approved" and _auto_refund_possible(p) and not p.refunded_cents, "auto_partial": p.status == "approved" and _auto_partial_possible(p),
                  "refunds": [{"amount": r.amount_cents // 100, "kind": r.kind, "status": r.status, "reason": r.reason, "note": r.note or "", "at": r.created_at.strftime("%Y-%m-%d %H:%M"), "id": r.id, "response": (r.wompi_response or "")[:600]} for r in refunds.get(p.id, [])],
                  "created_at": to_local(p.created_at).strftime("%Y-%m-%d %H:%M:%S"), "confirmed_at": to_local(p.confirmed_at).strftime("%Y-%m-%d %H:%M:%S") if p.confirmed_at else "",
                  "applied": json.loads(p.breakdown_json or "{}").get("applied", [])} for p in rows]}


# ------------------------------------------------------------------ base para pre-llenar y invitaciones
@router.post("/events/{event_id}/forms/{form_id}/people")
async def upload_people(event_id: int, form_id: int, file: UploadFile = File(...), db: Session = Depends(get_db), staff: StaffUser = Depends(STAFF)):
    """Sube (Excel/CSV) la base que pre-llenará SOLO este formulario (fuente «base subida aparte»); reemplaza la anterior."""
    from app.routers.api import _read_roster_rows, _ID_KEYS, _normalize_optional_key
    event = get_event_for_staff(event_id, db, staff)
    form = _get_form(db, event, form_id)
    header, rows = _read_roster_rows(file.filename, await file.read())
    id_col = next((k for k in header if formsvc.COLUMN_TO_KEY.get(k) == "id"), None)
    if not id_col:
        raise HTTPException(status_code=400, detail="El archivo necesita una columna de cédula (id / cédula)")
    db.query(FormPerson).filter(FormPerson.form_id == form.id).delete()
    seen, count = set(), 0
    for row in rows:
        person = {}
        for col, value in row.items():
            key = formsvc.COLUMN_TO_KEY.get(col) or _normalize_optional_key(col)
            if key and str(value).strip():
                person[key] = str(value).strip()
        pid = person.get("id")
        if pid and pid not in seen:
            seen.add(pid)
            db.add(FormPerson(form_id=form.id, person_id=pid, data_json=json.dumps(person)))
            count += 1
    settings = formsvc.get_settings(form)
    settings["prefill"]["source"] = "upload"
    form.settings_json = json.dumps(settings)
    db.commit()
    return {"people": count, "message": f"{count} personas cargadas para pre-llenar este formulario"}


@router.get("/events/{event_id}/forms/{form_id}/invites")
async def invites_status(event_id: int, form_id: int, db: Session = Depends(get_db), staff: StaffUser = Depends(STAFF)):
    event = get_event_for_staff(event_id, db, staff)
    form = _get_form(db, event, form_id)
    rows = db.query(FormInvite).filter_by(form_id=form.id).all()
    return {"total": len(rows), "sent": sum(1 for r in rows if r.sent_at), "used": sum(1 for r in rows if r.used_at), "source_people": len(formsvc.source_people(db, form))}


@router.post("/events/{event_id}/forms/{form_id}/invites/send")
async def send_invites(event_id: int, form_id: int, data: dict, request: Request, db: Session = Depends(get_db), staff: StaffUser = Depends(STAFF)):
    """Genera el enlace personal de cada persona de la fuente y lo manda por correo (a quien tenga correo) en segundo
    plano, con avance (`GET /api/bulk_jobs/{id}`). `only_unsent` (por defecto sí) no reenvía a quien ya lo recibió."""
    event = get_event_for_staff(event_id, db, staff)
    form = _get_form(db, event, form_id)
    if formsvc.get_settings(form)["prefill"]["mode"] != "invite":
        raise HTTPException(status_code=400, detail="Este formulario no usa enlaces personalizados por invitado (cámbialo en su configuración)")
    if bulk_jobs.active_job(db, event_id):
        raise HTTPException(status_code=409, detail="Ya hay un proceso en curso para este evento — espera a que termine")
    only_unsent = data.get("only_unsent", True)
    base = _base_url(request)
    job_id = bulk_jobs.create_job(db, event_id, staff.id)
    form_id_, slug, ev_id, ev_name, name = form.id, form.slug, event.id, event.name, form.name

    def work(job_db, reporter):
        f = job_db.query(WebForm).filter(WebForm.id == form_id_).first()
        people = formsvc.source_people(job_db, f)
        total = len(people)
        sent = skipped = failed = 0
        problems = []
        for i, p in enumerate(people, start=1):
            reporter(f"Enviando invitaciones ({i} de {total})", i, total)
            pid = str(p.get("id") or "")
            inv = job_db.query(FormInvite).filter_by(form_id=form_id_, person_id=pid).first()
            if not inv:
                inv = FormInvite(form_id=form_id_, person_id=pid, token=secrets.token_urlsafe(18), email=(p.get("email") or None))
                job_db.add(inv)
                job_db.commit()
            if not p.get("email"):
                skipped += 1
                continue
            if only_unsent and inv.sent_at:
                skipped += 1
                continue
            link = f"{base}/f/{ev_id}/{slug}?i={inv.token}"
            res = send_mail(p["email"], f"Tu inscripción — {ev_name}",
                            f"Hola {p.get('first_name') or ''}, te invitamos a completar tu inscripción a {ev_name} ({name}).\n\nAbre tu enlace personal (ya trae tus datos):\n{link}\n\nEs personal: no lo compartas.")
            if res.get("sent"):
                inv.sent_at = datetime.utcnow()
                job_db.commit()
                sent += 1
            else:
                failed += 1
                if len(problems) < 10:
                    problems.append(f"{p['email']}: {res.get('detail')}")
        return {"message": f"Invitaciones: {sent} enviadas, {skipped} omitidas (sin correo o ya enviadas), {failed} fallidas.", "sent": sent, "skipped": skipped, "failed": failed, "problems": problems}

    bulk_jobs.start(job_id, work)
    return {"background": True, "job_id": job_id}


# ------------------------------------------------------------------ analítica (módulo compartido, app/analytics.py)
def form_dashboard(db: Session, form: WebForm, include_tests: bool = False, net_fees: bool = False) -> dict:
    from collections import Counter
    from app import analytics as an

    design = formsvc.get_design(form)
    q_sub = db.query(FormSubmission).filter(FormSubmission.form_id == form.id, FormSubmission.status == "confirmed")
    q_evt = db.query(FormEvent).filter(FormEvent.form_id == form.id)
    if not include_tests:
        q_sub, q_evt = q_sub.filter(FormSubmission.is_test == False), q_evt.filter(FormEvent.is_test == False)  # noqa: E712
    subs = q_sub.order_by(FormSubmission.id).all()
    events = q_evt.all()
    views = {e.sid for e in events if e.kind == "view"}
    starts = {e.sid for e in events if e.kind == "start"}
    sent = {e.sid for e in events if e.kind == "submit"}
    abandoned = starts - sent
    durations = [(s.created_at - s.started_at).total_seconds() for s in subs if s.started_at and s.created_at >= s.started_at]

    state = formsvc.public_state(db, form)
    reason = {"cupo_lleno": "Cupo lleno", "cerrado": "Cerrado (manual o por fecha)", "finalizado": "Finalizado", "activo": "Abierto", "pruebas": "En pruebas"}[state]
    kpis = [
        an.kpi("Inscripciones", len(subs), "sin las de prueba" if not include_tests else "incluye pruebas", "good"),
        an.kpi("Estado", reason, "ahora mismo"),
        an.kpi("Visitas", len(views), "personas que abrieron el formulario"),
        an.kpi("Empezaron a llenarlo", len(starts), an.pct(len(starts), len(views)) + " de las visitas"),
        an.kpi("Lo enviaron", len(sent), an.pct(len(sent), len(views)) + " de las visitas (conversión)", "good"),
        an.kpi("Lo abandonaron", len(abandoned), an.pct(len(abandoned), len(starts)) + " de quienes empezaron", "warn" if abandoned else ""),
    ]
    if durations:
        durations.sort()
        kpis.append(an.kpi("Tiempo promedio", an.fmt_duration(sum(durations) / len(durations)), f"mediana {an.fmt_duration(durations[len(durations) // 2])}"))
    if form.capacity:
        kpis.append(an.kpi("Cupo", f"{len(subs) if include_tests else formsvc.real_submissions(db, form).count()} / {form.capacity}", "inscripciones reales / cupo", "warn" if formsvc.is_full(db, form) else ""))
    if subs:
        kpis.append(an.kpi("Última inscripción", an.fmt_local(subs[-1].created_at), "hora local"))

    charts = [an.time_series("timeline", "Inscripciones en el tiempo", [s.created_at for s in subs], label="Inscripciones"),
              an.hour_histogram("by_hour", "Inscripciones por hora del día", [s.created_at for s in subs], label="Inscripciones"),
              an.counter_chart("by_source", "¿De dónde llegaron? (enlaces con utm_source)", [s.source or "Directo" for s in subs], kind="pie", label="Inscripciones")]

    if formlib.payment_field(design):
        _payment_analytics(db, form, include_tests, kpis, charts, net_fees)

    # Campos con más y con menos respuesta (sobre las inscripciones, contando solo las que tenían el campo visible: aproximación por valor presente)
    if subs:
        answered = Counter()
        for s in subs:
            for fid, v in json.loads(s.data_json).items():
                if v not in (None, "", []):
                    answered[fid] += 1
        cols = _column_defs(design)
        rates = [(label, round(100 * answered.get(fid, 0) / len(subs))) for fid, label, _ in cols]
        if rates:
            rates.sort(key=lambda x: -x[1])
            charts.append(an.chart("response_rate", "Campos con más y menos respuesta (% de inscripciones que lo llenaron)", "bar", [r[0] for r in rates], [{"label": "% que respondió", "data": [r[1] for r in rates]}], horizontal=True,
                                   note="Un campo condicional solo lo ven algunos, por eso puede aparecer con menos respuestas."))
        # Distribución de cada campo marcado «genera estadísticas»
        for fid, label, kind in cols:
            f = design["fields"][fid]
            if not f.get("stats"):
                continue
            values = []
            for s in subs:
                v = json.loads(s.data_json).get(fid)
                if isinstance(v, list):
                    values += v
                elif v is True:
                    values.append("Sí")
                elif v is False:
                    values.append("No")
                elif isinstance(v, str) and kind in ("select", "radio", "checkbox", "number", "text_short", "email"):
                    values.append(v if kind != "email" else v.split("@")[-1])       # correo: por dominio, no por persona
            c = an.counter_chart(f"field_{fid}", f"{label}", values, kind="pie" if kind in ("select", "radio", "checkbox") else "bar", label="Inscripciones")
            if c:
                charts.append(c)
    return an.dashboard(f"Formulario: {form.name}", kpis, [c for c in charts if c])


def _payment_analytics(db: Session, form: WebForm, include_tests: bool, kpis: list, charts: list, net_fees: bool = False) -> None:
    """Ingresos del formulario (solo los pagos aprobados cuentan como ingreso) y cómo se resolvieron los intentos."""
    from collections import Counter, defaultdict
    from app import analytics as an

    q = db.query(FormPayment).filter(FormPayment.form_id == form.id)
    if not include_tests:
        q = q.filter(FormPayment.is_test == False)  # noqa: E712
    pays = q.order_by(FormPayment.id).all()
    if not pays:
        return
    pesos = lambda n: "$" + f"{n:,}".replace(",", ".")
    approved = [p for p in pays if p.status in ("approved", "refunded")]
    total = sum(p.amount_cents for p in approved) // 100
    refunded = sum(p.refunded_cents or 0 for p in approved) // 100
    failed = [p for p in pays if p.status in ("declined", "error", "voided")]
    abandoned = [p for p in pays if _pay_label(p) == "abandoned"]
    fees = sum(_fee_cop(p) for p in approved)
    if net_fees:
        kpis += [an.kpi("Comisión Wompi (estimada)", pesos(fees), "2,65 % + $700 + IVA por pago (QR 1 %); estimación según tu plan", "warn"),
                 an.kpi("Ingresos netos de comisión", pesos(total - refunded - fees), "aprobados − reembolsos − comisión estimada", "good")]
    kpis += [an.kpi("Ingresos (aprobados)", pesos(total), f"{len(approved)} pago(s) aprobado(s)", "good"),
             an.kpi("Pago promedio", pesos(total // len(approved) if approved else 0), "por inscripción pagada"),
             an.kpi("Pagos no completados", len(failed) + len(abandoned), f"{len(failed)} rechazados · {len(abandoned)} abandonados", "warn" if failed or abandoned else "")]
    if refunded:
        kpis += [an.kpi("Reembolsos", pesos(refunded), f"{sum(1 for p in approved if p.refunded_cents)} pago(s) reembolsado(s)", "warn"),
                 an.kpi("Ingresos netos", pesos(total - refunded), "aprobados menos reembolsos", "good")]
    orphans = sum(1 for p in pays if p.status == "approved" and not p.submission_id)
    if orphans:
        kpis.append(an.kpi("Pagos sin inscripción", orphans, "aprobados pero sin inscripción: revisar en la lista de pagos", "warn"))
    per_day = defaultdict(int)
    for p in approved:
        per_day[to_local(p.confirmed_at or p.created_at).strftime("%Y-%m-%d")] += (p.amount_cents // 100 - _fee_cop(p)) if net_fees else p.amount_cents // 100
    if per_day:
        days = sorted(per_day)
        charts.append(an.chart("revenue_by_day", "Ingresos por día (COP" + (", netos de comisión estimada)" if net_fees else ")"), "bar", days, [{"label": "Ingresos netos de comisión" if net_fees else "Ingresos aprobados", "data": [per_day[d] for d in days]}]))
    names = {"refunded": "Reembolsado", "approved": "Aprobado", "declined": "Rechazado", "error": "Error", "voided": "Anulado", "pending": "En espera", "abandoned": "Abandonado"}
    charts.append(an.counter_chart("payment_status", "Resultado de los intentos de pago", [names.get(_pay_label(p), _pay_label(p)) for p in pays], kind="pie", label="Pagos"))
    applied = Counter()
    for p in approved:
        for a in json.loads(p.breakdown_json or "{}").get("applied", []):
            applied[f"{a['label']} ({a['effect']})"] += 1
    if applied:
        charts.append(an.counter_chart("price_rules", "Reglas de precio y descuentos aplicados", list(applied.elements()), kind="bar", label="Pagos"))


@router.get("/events/{event_id}/forms/{form_id}/analytics")
async def form_analytics(event_id: int, form_id: int, include_tests: bool = False, net_fees: bool = False, db: Session = Depends(get_db), staff: StaffUser = Depends(STAFF)):
    event = get_event_for_staff(event_id, db, staff)
    return form_dashboard(db, _get_form(db, event, form_id), include_tests, net_fees)
