"""Control de Áreas y Control de Inventario (reunión 2026-09-21, ítem 9) — funcionalidad genérica por
evento, cada módulo con su switch en Parámetros del Evento.

Áreas: zonas del evento; entrada/salida por cédula/QR (el QR trae el ID) o reconocimiento facial, con el
mismo patrón que el Registro: con "Modo autoregistro" ON el match guarda el movimiento solo; OFF pide
confirmar ("Acreditar"). Por zona, "permitir reingresos"; apagado = solo una entrada por persona (alerta roja).

Inventario: ítems con cantidad inicial; se entregan uno o varios a la vez (combos) a una persona y el
disponible se descuenta solo (disponible = inicial - entregado)."""
import io
import uuid
from datetime import datetime

import pandas as pd
from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl.styles import Font, PatternFill
from sqlalchemy import func
from sqlalchemy.orm import Session

from app.auth import get_event_for_staff, require_event_in_progress, require_role_excluding
from app.biometrics import BiometricEngine
from app.database import get_db
from app.timeutil import to_local
from app.models import (
    AreaMovement, EventArea, EventAttendee, InventoryDelivery, InventoryItem, StaffUser, User,
)

router = APIRouter()

CONFIG_ROLE = require_role_excluding("coordinador", ("comercial",))  # define zonas/ítems y baja reportes
OPERATE_ROLE = require_role_excluding("digitador", ("comercial",))  # registra movimientos / entregas


def delete_event_modules(db: Session, event_id: int) -> None:
    """Borra los datos de Áreas/Inventario de un evento (lo usa delete_event)."""
    db.query(AreaMovement).filter(AreaMovement.event_id == event_id).delete()
    db.query(EventArea).filter(EventArea.event_id == event_id).delete()
    db.query(InventoryDelivery).filter(InventoryDelivery.event_id == event_id).delete()
    db.query(InventoryItem).filter(InventoryItem.event_id == event_id).delete()


@router.put("/events/{event_id}/modules")
async def set_modules(event_id: int, data: dict, db: Session = Depends(get_db), staff: StaffUser = Depends(CONFIG_ROLE)):
    """Switches por evento: `areas_enabled` / `inventory_enabled` (solo se cambian los que vengan)."""
    event = get_event_for_staff(event_id, db, staff)
    if "areas_enabled" in data:
        event.areas_enabled = bool(data["areas_enabled"])
    if "inventory_enabled" in data:
        event.inventory_enabled = bool(data["inventory_enabled"])
    db.commit()
    return {"areas_enabled": event.areas_enabled, "inventory_enabled": event.inventory_enabled}


def _person_data(user: User) -> dict:
    return {"id": user.id, "first_name": user.first_name, "last_name": user.last_name, "entity": user.entity, "opt_1": user.opt_1}


def _find_person(db: Session, event, cedula: str) -> User:
    """La persona debe estar en la base de ESTE evento (no cualquier persona del cliente)."""
    cedula = (cedula or "").strip()
    user = db.query(User).filter(User.id == cedula, User.tenant_id == event.tenant_id).first()
    in_event = user and db.query(EventAttendee).filter_by(event_id=event.id, user_id=cedula).first()
    if not in_event:
        raise HTTPException(status_code=404, detail="Esa persona no está en la base de este evento")
    return user


# ---------------------------------------------------------------------------------------- Áreas

def _area_json(db: Session, a: EventArea) -> dict:
    last_by_user = {}
    for m in db.query(AreaMovement).filter(AreaMovement.area_id == a.id).order_by(AreaMovement.timestamp, AreaMovement.id):
        last_by_user[m.user_id] = m.direction
    return {
        "id": a.id, "name": a.name, "allow_reentry": a.allow_reentry,
        "inside": sum(1 for d in last_by_user.values() if d == "in"), "visited": len(last_by_user),
    }


def _get_area(db: Session, event, area_id: int) -> EventArea:
    area = db.query(EventArea).filter(EventArea.id == area_id, EventArea.event_id == event.id).first()
    if not area:
        raise HTTPException(status_code=404, detail="Zona no encontrada")
    return area


@router.get("/events/{event_id}/areas")
async def list_areas(event_id: int, db: Session = Depends(get_db), staff: StaffUser = Depends(OPERATE_ROLE)):
    event = get_event_for_staff(event_id, db, staff)
    return [_area_json(db, a) for a in db.query(EventArea).filter(EventArea.event_id == event.id).order_by(EventArea.id)]


@router.post("/events/{event_id}/areas")
async def create_area(event_id: int, data: dict, db: Session = Depends(get_db), staff: StaffUser = Depends(CONFIG_ROLE)):
    event = get_event_for_staff(event_id, db, staff)
    name = str(data.get("name") or "").strip()
    if not name:
        raise HTTPException(status_code=400, detail="La zona necesita un nombre")
    if db.query(EventArea).filter(EventArea.event_id == event.id, func.lower(EventArea.name) == name.lower()).first():
        raise HTTPException(status_code=400, detail="Ya existe una zona con ese nombre en este evento")
    area = EventArea(event_id=event.id, name=name, allow_reentry=bool(data.get("allow_reentry", True)), created_at=datetime.utcnow())
    db.add(area)
    db.commit()
    return _area_json(db, area)


@router.patch("/events/{event_id}/areas/{area_id}")
async def update_area(event_id: int, area_id: int, data: dict, db: Session = Depends(get_db), staff: StaffUser = Depends(CONFIG_ROLE)):
    area = _get_area(db, get_event_for_staff(event_id, db, staff), area_id)
    if "name" in data and str(data["name"]).strip():
        area.name = str(data["name"]).strip()
    if "allow_reentry" in data:
        area.allow_reentry = bool(data["allow_reentry"])
    db.commit()
    return _area_json(db, area)


@router.delete("/events/{event_id}/areas/{area_id}")
async def delete_area(event_id: int, area_id: int, db: Session = Depends(get_db), staff: StaffUser = Depends(CONFIG_ROLE)):
    area = _get_area(db, get_event_for_staff(event_id, db, staff), area_id)
    db.query(AreaMovement).filter(AreaMovement.area_id == area.id).delete()
    db.delete(area)
    db.commit()
    return {"message": "Zona eliminada (con sus movimientos)"}


def _apply_movement(db: Session, event, area: EventArea, user: User, direction: str, confirm: bool, method: str, staff: StaffUser) -> dict:
    """Reglas de una entrada/salida. Devuelve el resultado (nunca lanza por reglas de negocio):
    OK (guardado) | FOUND_PENDING (autoregistro apagado y sin confirmar) | ALREADY_ENTERED (alerta roja:
    reingreso no permitido) | NOT_INSIDE (salida de alguien que no está adentro)."""
    last = db.query(AreaMovement).filter(AreaMovement.area_id == area.id, AreaMovement.user_id == user.id) \
        .order_by(AreaMovement.timestamp.desc(), AreaMovement.id.desc()).first()
    if direction == "auto":
        direction = "out" if (last and last.direction == "in") else "in"
    if direction not in ("in", "out"):
        raise HTTPException(status_code=400, detail="direction debe ser in, out o auto")
    data = _person_data(user)

    if direction == "in" and not area.allow_reentry:
        already = db.query(AreaMovement).filter(
            AreaMovement.area_id == area.id, AreaMovement.user_id == user.id, AreaMovement.direction == "in"
        ).first()
        if already:
            return {"result": "ALREADY_ENTERED", "direction": "in", "area": area.name, "data": data,
                    "details": f"{user.first_name} {user.last_name} YA había entrado a «{area.name}» y esta zona no permite reingresos"}
    if direction == "out" and not (last and last.direction == "in"):
        return {"result": "NOT_INSIDE", "direction": "out", "area": area.name, "data": data,
                "details": f"{user.first_name} {user.last_name} no figura adentro de «{area.name}»"}

    if not event.auto_register and not confirm:
        return {"result": "FOUND_PENDING", "direction": direction, "area": area.name, "data": data}

    now = datetime.utcnow()
    db.add(AreaMovement(
        event_id=event.id, area_id=area.id, user_id=user.id, tenant_id=event.tenant_id, direction=direction,
        method=method, timestamp=now, registered_by_staff_id=staff.id,
    ))
    db.commit()
    return {"result": "OK", "direction": direction, "area": area.name, "data": data, "time": now.isoformat()}


@router.post("/events/{event_id}/areas/{area_id}/movement")
async def area_movement(event_id: int, area_id: int, data: dict, db: Session = Depends(get_db), staff: StaffUser = Depends(OPERATE_ROLE)):
    """Entrada/salida por cédula o QR (el QR trae el ID). `direction`: in | out | auto (alterna)."""
    event = get_event_for_staff(event_id, db, staff)
    require_event_in_progress(event)
    area = _get_area(db, event, area_id)
    user = _find_person(db, event, str(data.get("cedula") or ""))
    return _apply_movement(db, event, area, user, data.get("direction", "auto"), bool(data.get("confirm")), "qr" if data.get("method") == "qr" else "cedula", staff)


@router.post("/events/{event_id}/areas/{area_id}/movement-face")
async def area_movement_face(
    event_id: int, area_id: int, file: UploadFile = File(...), direction: str = Form("auto"), confirm: bool = Form(False),
    db: Session = Depends(get_db), staff: StaffUser = Depends(OPERATE_ROLE),
):
    """Entrada/salida por reconocimiento facial (solo eventos con biometría)."""
    event = get_event_for_staff(event_id, db, staff)
    require_event_in_progress(event)
    area = _get_area(db, event, area_id)
    img_array = BiometricEngine.process_image_stream(await file.read())
    unknown = BiometricEngine.extract_encoding(img_array)
    if not unknown:
        return {"result": "NO", "details": "Rostro no detectado"}
    attendee_ids = {a.user_id for a in db.query(EventAttendee).filter(EventAttendee.event_id == event.id)}
    for user in db.query(User).filter(User.tenant_id == event.tenant_id, User.id.in_(attendee_ids)).all():
        known = user.get_encoding()
        if known and BiometricEngine.compare(known, unknown):
            return _apply_movement(db, event, area, user, direction, confirm, "facial", staff)
    return {"result": "NO", "details": "Ninguna persona de este evento coincide"}


@router.get("/events/{event_id}/areas/{area_id}/presence")
async def area_presence(event_id: int, area_id: int, db: Session = Depends(get_db), staff: StaffUser = Depends(OPERATE_ROLE)):
    """Última situación de cada persona en la zona: verde = entró ('in'), rojo = salió ('out')."""
    event = get_event_for_staff(event_id, db, staff)
    area = _get_area(db, event, area_id)
    last = {}
    for m in db.query(AreaMovement).filter(AreaMovement.area_id == area.id).order_by(AreaMovement.timestamp, AreaMovement.id):
        last[m.user_id] = m
    users = {u.id: u for u in db.query(User).filter(User.tenant_id == event.tenant_id, User.id.in_(list(last.keys()) or [""])).all()}
    rows = [{
        "id": uid, "name": f"{users[uid].first_name or ''} {users[uid].last_name or ''}".strip(),
        "direction": m.direction, "time": m.timestamp.isoformat(),
    } for uid, m in last.items() if uid in users]
    rows.sort(key=lambda r: r["time"], reverse=True)
    return rows


def _xlsx_response(sheets: dict, filename: str, widths: dict = None) -> StreamingResponse:
    out = io.BytesIO()
    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        for name, df in sheets.items():
            df.to_excel(writer, index=False, sheet_name=name)
            ws = writer.sheets[name]
            for cell in ws[1]:
                cell.fill = PatternFill(start_color="0A0E2E", end_color="0A0E2E", fill_type="solid")
                cell.font = Font(color="FFFFFF", bold=True)
            for idx, col in enumerate(df.columns, start=1):
                longest = max([len(str(col))] + [len(str(v)) for v in df[col]])
                ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = min(longest + 3, 45)
            ws.auto_filter.ref = ws.dimensions
            ws.freeze_panes = "A2"
    out.seek(0)
    return StreamingResponse(
        out, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/events/{event_id}/areas-report")
async def areas_report(event_id: int, db: Session = Depends(get_db), staff: StaffUser = Depends(CONFIG_ROLE)):
    """Todos los movimientos con hora EXACTA (con segundos) y zona."""
    event = get_event_for_staff(event_id, db, staff)
    areas = {a.id: a.name for a in db.query(EventArea).filter(EventArea.event_id == event.id)}
    users = {u.id: u for u in db.query(User).filter(User.tenant_id == event.tenant_id)}
    staff_names = {s.id: (s.full_name or s.username) for s in db.query(StaffUser)}
    rows = []
    for m in db.query(AreaMovement).filter(AreaMovement.event_id == event.id).order_by(AreaMovement.timestamp, AreaMovement.id):
        u = users.get(m.user_id)
        rows.append([
            to_local(m.timestamp).strftime("%Y-%m-%d"), to_local(m.timestamp).strftime("%H:%M:%S"), areas.get(m.area_id, ""), m.user_id,
            f"{u.first_name or ''} {u.last_name or ''}".strip() if u else "", "Entrada" if m.direction == "in" else "Salida",
            m.method or "", staff_names.get(m.registered_by_staff_id, ""),
        ])
    df = pd.DataFrame(rows, columns=["Fecha", "Hora", "Zona", "Cédula", "Persona", "Movimiento", "Método", "Registrado por"])
    return _xlsx_response({"Movimientos por zona": df}, f"Areas_{event.event_code}.xlsx")


# ------------------------------------------------------------------------------------ Inventario

def _delivered_by_item(db: Session, event_id: int) -> dict:
    rows = db.query(InventoryDelivery.item_id, func.sum(InventoryDelivery.qty)).filter(InventoryDelivery.event_id == event_id).group_by(InventoryDelivery.item_id).all()
    return {item_id: int(total or 0) for item_id, total in rows}


def _item_json(item: InventoryItem, delivered: int) -> dict:
    return {
        "id": item.id, "name": item.name, "initial_qty": item.initial_qty, "allow_multiple": item.allow_multiple,
        "delivered": delivered, "available": item.initial_qty - delivered,
    }


def _get_item(db: Session, event, item_id: int) -> InventoryItem:
    item = db.query(InventoryItem).filter(InventoryItem.id == item_id, InventoryItem.event_id == event.id).first()
    if not item:
        raise HTTPException(status_code=404, detail="Ítem no encontrado")
    return item


def _qty(value, minimum: int = 0) -> int:
    try:
        n = int(value)
    except (TypeError, ValueError):
        raise HTTPException(status_code=400, detail="La cantidad debe ser un número entero")
    if n < minimum:
        raise HTTPException(status_code=400, detail=f"La cantidad debe ser {minimum} o más")
    return n


@router.get("/events/{event_id}/inventory")
async def list_inventory(event_id: int, db: Session = Depends(get_db), staff: StaffUser = Depends(OPERATE_ROLE)):
    event = get_event_for_staff(event_id, db, staff)
    delivered = _delivered_by_item(db, event.id)
    return [_item_json(i, delivered.get(i.id, 0)) for i in db.query(InventoryItem).filter(InventoryItem.event_id == event.id).order_by(InventoryItem.id)]


@router.post("/events/{event_id}/inventory")
async def create_item(event_id: int, data: dict, db: Session = Depends(get_db), staff: StaffUser = Depends(CONFIG_ROLE)):
    event = get_event_for_staff(event_id, db, staff)
    name = str(data.get("name") or "").strip()
    if not name:
        raise HTTPException(status_code=400, detail="El ítem necesita un nombre")
    if db.query(InventoryItem).filter(InventoryItem.event_id == event.id, func.lower(InventoryItem.name) == name.lower()).first():
        raise HTTPException(status_code=400, detail="Ya existe un ítem con ese nombre en este evento")
    item = InventoryItem(event_id=event.id, name=name, initial_qty=_qty(data.get("initial_qty", 0)), allow_multiple=bool(data.get("allow_multiple")), created_at=datetime.utcnow())
    db.add(item)
    db.commit()
    return _item_json(item, 0)


@router.patch("/events/{event_id}/inventory/{item_id}")
async def update_item(event_id: int, item_id: int, data: dict, db: Session = Depends(get_db), staff: StaffUser = Depends(CONFIG_ROLE)):
    event = get_event_for_staff(event_id, db, staff)
    item = _get_item(db, event, item_id)
    delivered = _delivered_by_item(db, event.id).get(item.id, 0)
    if "name" in data and str(data["name"]).strip():
        item.name = str(data["name"]).strip()
    if "initial_qty" in data:
        new_qty = _qty(data["initial_qty"])
        if new_qty < delivered:
            raise HTTPException(status_code=400, detail=f"Ya se entregaron {delivered}: la cantidad inicial no puede ser menor")
        item.initial_qty = new_qty
    if "allow_multiple" in data:
        item.allow_multiple = bool(data["allow_multiple"])
    db.commit()
    return _item_json(item, delivered)


@router.delete("/events/{event_id}/inventory/{item_id}")
async def delete_item(event_id: int, item_id: int, db: Session = Depends(get_db), staff: StaffUser = Depends(CONFIG_ROLE)):
    event = get_event_for_staff(event_id, db, staff)
    item = _get_item(db, event, item_id)
    if _delivered_by_item(db, event.id).get(item.id, 0):
        raise HTTPException(status_code=400, detail="Este ítem ya tiene entregas registradas — no se puede borrar (ajusta su cantidad si hace falta)")
    db.delete(item)
    db.commit()
    return {"message": "Ítem eliminado"}


@router.post("/events/{event_id}/inventory/deliver")
async def deliver(event_id: int, data: dict, db: Session = Depends(get_db), staff: StaffUser = Depends(OPERATE_ROLE)):
    """Entrega uno o varios ítems (combo) a una persona: `{cedula, items: [{item_id, qty}]}`. Todo o
    nada: si algún ítem no alcanza, no se entrega nada. El disponible se descuenta solo."""
    event = get_event_for_staff(event_id, db, staff)
    require_event_in_progress(event)
    user = _find_person(db, event, str(data.get("cedula") or ""))
    lines = data.get("items") or []
    if not lines:
        raise HTTPException(status_code=400, detail="Elige al menos un ítem para entregar")

    wanted = {}
    for line in lines:
        item_id = int(line.get("item_id"))
        wanted[item_id] = wanted.get(item_id, 0) + _qty(line.get("qty", 1), minimum=1)

    # Con lock de fila: dos digitadores entregando el último ítem a la vez no lo pueden vender dos veces.
    items = {i.id: i for i in db.query(InventoryItem).filter(InventoryItem.event_id == event.id, InventoryItem.id.in_(list(wanted))).with_for_update().all()}
    delivered = _delivered_by_item(db, event.id)
    for item_id, qty in wanted.items():
        item = items.get(item_id)
        if not item:
            raise HTTPException(status_code=404, detail="Uno de los ítems no existe en este evento")
        if qty > 1 and not item.allow_multiple:
            raise HTTPException(status_code=400, detail=f"«{item.name}» solo se entrega de a 1 unidad por vez")
        available = item.initial_qty - delivered.get(item_id, 0)
        if qty > available:
            raise HTTPException(status_code=400, detail=f"No alcanza «{item.name}»: quedan {available} y se pidieron {qty}")

    batch, now = uuid.uuid4().hex, datetime.utcnow()
    for item_id, qty in wanted.items():
        db.add(InventoryDelivery(
            event_id=event.id, item_id=item_id, batch_id=batch, user_id=user.id, tenant_id=event.tenant_id,
            qty=qty, delivered_at=now, delivered_by_staff_id=staff.id,
        ))
    db.commit()
    delivered = _delivered_by_item(db, event.id)
    return {
        "result": "OK", "data": _person_data(user),
        "delivered": [{"name": items[i].name, "qty": q} for i, q in wanted.items()],
        "stock": [_item_json(items[i], delivered.get(i, 0)) for i in wanted],
    }


@router.get("/events/{event_id}/inventory/deliveries")
async def recent_deliveries(event_id: int, db: Session = Depends(get_db), staff: StaffUser = Depends(OPERATE_ROLE)):
    event = get_event_for_staff(event_id, db, staff)
    items = {i.id: i.name for i in db.query(InventoryItem).filter(InventoryItem.event_id == event.id)}
    users = {u.id: u for u in db.query(User).filter(User.tenant_id == event.tenant_id)}
    rows = db.query(InventoryDelivery).filter(InventoryDelivery.event_id == event.id).order_by(InventoryDelivery.delivered_at.desc(), InventoryDelivery.id.desc()).limit(60).all()
    return [{
        "time": r.delivered_at.isoformat(), "user_id": r.user_id,
        "name": f"{users[r.user_id].first_name or ''} {users[r.user_id].last_name or ''}".strip() if r.user_id in users else "",
        "item": items.get(r.item_id, ""), "qty": r.qty,
    } for r in rows]


@router.get("/events/{event_id}/inventory-report")
async def inventory_report(event_id: int, db: Session = Depends(get_db), staff: StaffUser = Depends(CONFIG_ROLE)):
    """Hoja 'Entregas' (hora, a quién, qué se le entregó) + hoja 'Inventario' (inicial / entregado / disponible)."""
    event = get_event_for_staff(event_id, db, staff)
    items = {i.id: i for i in db.query(InventoryItem).filter(InventoryItem.event_id == event.id).order_by(InventoryItem.id)}
    users = {u.id: u for u in db.query(User).filter(User.tenant_id == event.tenant_id)}
    staff_names = {s.id: (s.full_name or s.username) for s in db.query(StaffUser)}
    deliveries = []
    for r in db.query(InventoryDelivery).filter(InventoryDelivery.event_id == event.id).order_by(InventoryDelivery.delivered_at, InventoryDelivery.id):
        u = users.get(r.user_id)
        deliveries.append([
            to_local(r.delivered_at).strftime("%Y-%m-%d"), to_local(r.delivered_at).strftime("%H:%M:%S"), r.user_id,
            f"{u.first_name or ''} {u.last_name or ''}".strip() if u else "", items[r.item_id].name if r.item_id in items else "",
            r.qty, staff_names.get(r.delivered_by_staff_id, ""),
        ])
    delivered = _delivered_by_item(db, event.id)
    stock = [[i.name, i.initial_qty, delivered.get(i.id, 0), i.initial_qty - delivered.get(i.id, 0)] for i in items.values()]
    return _xlsx_response({
        "Entregas": pd.DataFrame(deliveries, columns=["Fecha", "Hora", "Cédula", "Persona", "Ítem", "Cantidad", "Entregado por"]),
        "Inventario": pd.DataFrame(stock, columns=["Ítem", "Cantidad inicial", "Entregado", "Disponible"]),
    }, f"Inventario_{event.event_code}.xlsx")
