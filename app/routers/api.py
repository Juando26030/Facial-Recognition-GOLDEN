import os
import re
import json
from typing import Optional
import zipfile
import tempfile
import csv
import io
import face_recognition
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from PIL import Image
from fastapi import APIRouter, Depends, File, UploadFile, Form, HTTPException
from fastapi.responses import FileResponse
from sqlalchemy import func
from sqlalchemy.orm import Session
from app.database import get_db
from app.models import Event, User, AccessLog, EventAttendee, PrintLog, StaffUser
from app.biometrics import BiometricEngine
from app.reports import ReportManager
from app.auth import get_current_staff, get_event_for_staff, require_event_in_progress, require_role, require_role_excluding, require_role_or_client
from app.routers import parametros, signatures
from app.routers.events import _typo_match, _words


def _missing_required_fields(field_configs: list, values: dict) -> list:
    """Nombres (rótulos) de los campos marcados obligatorios en Parámetros del Evento que no
    vienen con valor en `values` — usado por manual_register y update_user (PATCH) para no dejar
    guardar un perfil incompleto. NO se usa en bulk_register (ver parametros.py, decisión de
    alcance). Un campo booleano "obligatorio" exige estar marcado (true), no solo presente."""
    missing = []
    for cfg in field_configs:
        if not cfg["required"]:
            continue
        value = values.get(cfg["key"])
        if cfg["field_type"] in ("boolean", "consent"):
            ok = str(value).strip().lower() in ("true", "1", "si", "sí")
        else:
            ok = value is not None and str(value).strip() != ""
        if not ok:
            missing.append(cfg["label"])
    return missing


def _read_roster_rows(filename: str, content: bytes):
    """Lee un roster en .csv o .xlsx y devuelve (header_columns, rows): filas normalizadas
    (claves en minúscula, sin espacios) sin importar el formato de origen, y un mapeo
    {header: 'A'|'B'|...} de en qué columna venía cada campo — así bulk_register puede señalar
    la celda exacta (ej. "B7") de un error, sea el archivo .xlsx o .csv."""
    filename = (filename or "").lower()
    if filename.endswith(".xlsx"):
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        headers = [str(h).strip().lower() if h is not None else "" for h in next(rows_iter, [])]
        header_columns = {h: get_column_letter(i + 1) for i, h in enumerate(headers) if h}
        rows = []
        for row in rows_iter:
            if row is None or all(cell is None for cell in row):
                continue
            rows.append({
                headers[i]: (str(cell).strip() if cell is not None else "")
                for i, cell in enumerate(row) if i < len(headers) and headers[i]
            })
        return header_columns, rows
    else:
        decoded = content.decode("utf-8-sig")
        delimiter = ";" if ";" in decoded.split("\n", 1)[0] else ","
        reader = csv.DictReader(io.StringIO(decoded), delimiter=delimiter)
        headers = [h.strip().lower() if h else "" for h in (reader.fieldnames or [])]
        header_columns = {h: get_column_letter(i + 1) for i, h in enumerate(headers) if h}
        rows = [
            {k.strip().lower() if k else "": (v or "").strip() for k, v in row.items() if k}
            for row in reader
        ]
        return header_columns, rows

MAX_OPTIONAL_FIELDS = 30
_OPTIONAL_FIELD_RE = re.compile(r"^opcional[\s_]*([0-9]{1,2})$")


def _normalize_optional_key(raw_key: str):
    """Reconoce cualquier variante de encabezado 'opcional 1'..'opcional 30' (con espacio, guion
    bajo, o pegado — 'opcional1') y la normaliza a la forma canónica 'opcional_N' que usamos para
    guardar en User.extra_fields y Event.optional_field_labels. None si no matchea o el número
    está fuera de rango."""
    if not raw_key:
        return None
    match = _OPTIONAL_FIELD_RE.match(raw_key.strip())
    if not match:
        return None
    n = int(match.group(1))
    return f"opcional_{n}" if 1 <= n <= MAX_OPTIONAL_FIELDS else None


def _optional_field_examples(rows: list, keys) -> dict:
    """Hasta 2 valores no vacíos (sin repetir) de cada campo opcional en `keys`, tal como vienen
    en el archivo — 2026-09-16, pedido explícito, para no tener que abrir el Excel a ver qué es
    'opcional_1'. `rows` ya viene con las claves originales del archivo (no normalizadas), así
    que hay que normalizar cada `raw_key` para saber a cuál "opcional_N" corresponde."""
    examples = {key: [] for key in keys}
    for row in rows:
        if all(len(v) >= 2 for v in examples.values()):
            break
        for raw_key, value in row.items():
            norm = _normalize_optional_key(raw_key)
            if norm not in examples:
                continue
            val = (value or "").strip()
            if val and val not in examples[norm] and len(examples[norm]) < 2:
                examples[norm].append(val)
    return examples


def _parse_extra_fields(raw: str) -> dict:
    """Parsea el JSON de extra_fields que manda el frontend (bulk_register por fila, o
    manual_register para una sola persona) y devuelve solo las claves 'opcional_N' válidas con
    valor no vacío, normalizadas."""
    if not raw:
        return {}
    try:
        provided = json.loads(raw)
    except (json.JSONDecodeError, TypeError):
        return {}
    extras = {}
    for raw_key, value in (provided or {}).items():
        norm = _normalize_optional_key(raw_key)
        val = str(value or "").strip()
        if norm and val:
            extras[norm] = val
    return extras


def _apply_optional_labels(event, used_keys: set, field_labels_raw: str) -> dict:
    """Fusiona field_labels_raw (JSON, ej. {"opcional_1": "Talla de camisa"}) dentro de
    Event.optional_field_labels — cualquier clave de used_keys que quede sin nombre después del
    merge recibe un rótulo por defecto ("Opcional N") para no dejar el diccionario incompleto.
    Compartido por bulk_register (carga de Excel) y manual_register (alta individual) para no
    duplicar el flujo de 'pregúntame a qué corresponde cada opcional'. No hace commit — quien
    llama decide cuándo."""
    existing = event.get_optional_labels()
    merged = dict(existing)
    try:
        provided = json.loads(field_labels_raw) if field_labels_raw else {}
    except (json.JSONDecodeError, TypeError):
        provided = {}
    for raw_key, label in (provided or {}).items():
        norm = _normalize_optional_key(raw_key)
        if norm and label and str(label).strip():
            merged[norm] = str(label).strip()
    for key in used_keys - set(merged.keys()):
        merged[key] = f"Opcional {key.split('_')[1]}"
    event.set_optional_labels(merged)
    return merged


_ID_KEYS = ('id', 'identificación', 'cedula', 'cédula')
_FLOAT_LOOKING_ID_RE = re.compile(r'^(\d+)\.0+$')
_EMAIL_RE = re.compile(r'^[^@\s]+@[^@\s]+\.[^@\s]+$')
_PHONE_RE = re.compile(r'^[\d\s+()\-]+$')


def _field_lookup(clean_row: dict, header_columns: dict, row_num: int, *keys: str):
    """Busca el primer valor no vacío entre varios encabezados alternativos (ej. 'nombres' o
    'nombre') y devuelve (valor, referencia_de_celda) — ej. ('Juan', 'B7'). Si ninguno tiene
    valor pero al menos uno de esos encabezados sí vino en el archivo, igual devuelve su celda
    (útil para señalar dónde falta el dato); si el archivo ni siquiera trae esa columna, la
    referencia es None y quien llama debe caer de vuelta a "fila N"."""
    for key in keys:
        val = clean_row.get(key, '')
        if val:
            col = header_columns.get(key)
            return val, (f"{col}{row_num}" if col else None)
    for key in keys:
        if key in header_columns:
            return '', f"{header_columns[key]}{row_num}"
    return '', None


def _extract_identificador(clean_row: dict, header_columns: dict, row_num: int):
    """Saca la cédula/ID de la fila (probando los encabezados alternativos de siempre) y de paso
    corrige el problema clásico de Excel de convertir una columna de cédulas en números, que
    entonces llegan como '1020304050.0' en vez de '1020304050'. Devuelve (identificador_limpio,
    celda, nota_o_None) — la nota, si existe, se le muestra al usuario explicando la corrección."""
    raw_val, cell_ref = _field_lookup(clean_row, header_columns, row_num, *_ID_KEYS)
    val = raw_val.strip()
    note = None
    match = _FLOAT_LOOKING_ID_RE.match(val)
    if match:
        corrected = match.group(1)
        where = cell_ref or f"fila {row_num}"
        note = f"ℹ️ Fila {row_num} ({where}): la cédula venía como '{val}' — Excel la convirtió a número. Se corrigió automáticamente a '{corrected}'."
        val = corrected
    return val, cell_ref, note


router = APIRouter()


def _known_faces_dir(tenant_id: str) -> str:
    path = os.path.join('data', tenant_id, 'known_people')
    os.makedirs(path, exist_ok=True)
    return path


def _upsert_attendee(db: Session, event_id: int, user_id: str, tenant_id: str, categories: Optional[list] = None) -> None:
    """Asegura que esta persona quede en la lista del evento, se haya precargado o no (ej. un
    'Nuevo' dado de alta sobre la marcha el día del evento) — ver EventAttendee en models.py.
    `categories` (ítem 14): si viene, fija las categorías de la persona EN este evento."""
    row = db.query(EventAttendee).filter_by(event_id=event_id, user_id=user_id).first()
    if not row:
        row = EventAttendee(event_id=event_id, user_id=user_id, tenant_id=tenant_id)
        db.add(row)
    if categories is not None:
        row.set_categories(categories)


def _clean_categories(event: Event, raw, auto_add: bool = False) -> list:
    """Categorías (ítem 14) de `raw` (lista o texto separado por , ; |) que existen en el evento,
    con la escritura oficial del evento. `auto_add=True` (carga de roster) agrega al evento las
    que no existían todavía, para que importar un Excel con categorías nuevas no las pierda."""
    if isinstance(raw, str):
        raw = re.split(r"[,;|]", raw)
    known = event.get_categories()
    by_lower = {c.lower(): c for c in known}
    result = []
    for name in (str(x).strip() for x in (raw or [])):
        if not name:
            continue
        canonical = by_lower.get(name.lower())
        if not canonical and auto_add:
            canonical = name
            known.append(name)
            by_lower[name.lower()] = name
        if canonical and canonical not in result:
            result.append(canonical)
    if auto_add:
        event.set_categories(known)
    return result


def _already_checked_in(db: Session, event_id: int, user_id: str) -> bool:
    """True si esta persona ya tiene al menos un AccessLog de ACREDITACIÓN real para ESTE evento
    — o sea, ya se acreditó por cualquier método (facial, cédula o alta manual). Excluye
    "Actualizado" (2026-09-17, mismo bug de fondo que el de get_all_users: ese record_type es solo
    una edición de perfil vía update_user, no una acreditación — contarlo acá haría que alguien
    que nunca se presentó, pero cuyo perfil ya se editó una vez, disparara el aviso de "ya
    registrado" en su primer check-in real, cuando en realidad es el primero)."""
    return db.query(AccessLog).filter(
        AccessLog.event_id == event_id, AccessLog.user_id == user_id, AccessLog.record_type != "Actualizado"
    ).first() is not None


def _duplicate_warning(db: Session, event_id: int, user: "User") -> dict:
    # times_registered (2026-09-16, pedido explícito): cuántas veces YA se acreditó esta persona
    # en este evento, para que el mensaje diga "ya se registró N veces" en vez de un genérico
    # "ya está registrada" — el operador decide con ese dato de más. Excluye "Actualizado"
    # (2026-09-17, mismo bug de fondo que _already_checked_in/get_all_users): son ediciones de
    # perfil, no acreditaciones — contarlas inflaría el número mostrado sin motivo real.
    times_registered = db.query(AccessLog).filter(
        AccessLog.event_id == event_id, AccessLog.user_id == user.id, AccessLog.record_type != "Actualizado"
    ).count()
    return {
        "result": "DUPLICADO", "details": "Esta persona ya había sido registrada en este evento",
        "times_registered": times_registered,
        "data": {
            "id": user.id, "first_name": user.first_name, "last_name": user.last_name,
            "role": user.role, "entity": user.entity, "phone": user.phone,
            "email": user.email, "opt_1": user.opt_1, "opt_2": user.opt_2
        },
    }


@router.get("/users")
async def get_all_users(
    event_id: int, db: Session = Depends(get_db), staff: StaffUser = Depends(get_current_staff)
):
    """Cualquier staff autenticado con acceso al evento puede VER el directorio (digitador y
    cliente incluidos) — get_event_for_staff abajo hace el chequeo real de autorización.
    El directorio es de ESTE evento: roster precargado (EventAttendee) unido con quien de hecho
    se presentó (AccessLog filtrado por event_id) — ya no todos los User del tenant sin distinguir
    entre eventos distintos del mismo cliente."""
    event = get_event_for_staff(event_id, db, staff)

    attendee_ids = {a.user_id for a in db.query(EventAttendee).filter(EventAttendee.event_id == event_id)}
    log_ids = {l.user_id for l in db.query(AccessLog).filter(AccessLog.event_id == event_id)}
    all_ids = attendee_ids | log_ids
    if not all_ids:
        return []

    users = db.query(User).filter(User.tenant_id == event.tenant_id, User.id.in_(all_ids)).all()
    categories_by_user = {a.user_id: a.get_categories() for a in db.query(EventAttendee).filter(EventAttendee.event_id == event_id)}
    event_logs = db.query(AccessLog).filter(AccessLog.event_id == event_id).all()
    logs_by_user = {}
    for log in event_logs:
        logs_by_user.setdefault(log.user_id, []).append(log)

    result = []
    for u in users:
        # Bug real (2026-09-17, reportado en QA): "Actualizado" es un log de EDICIÓN de perfil,
        # no de acreditación (ver update_user más abajo) — antes contaba igual que "Nuevo"/
        # "Existente" para decidir el estado, así que revertir a alguien a "No registrado" desde
        # el modal de Editar (que en el mismo clic, después del cambio de estado, también guarda
        # el resto del formulario vía update_user) volvía a dejarlo en "Registrado" de una,
        # porque ese mismo guardado crea un "Actualizado" nuevo apenas se borran los logs reales.
        real_logs = [log for log in logs_by_user.get(u.id, []) if log.record_type != "Actualizado"]
        status = "No registrado"
        if real_logs:
            status = "Nuevo" if any(log.record_type == "Nuevo" for log in real_logs) else "Registrado"

        result.append({
            "id": u.id, "first_name": u.first_name, "last_name": u.last_name,
            "role": u.role, "entity": u.entity, "phone": u.phone,
            "email": u.email, "opt_1": u.opt_1, "opt_2": u.opt_2, "status": status,
            # 2026-09-16: el modal de "Editar" del Directorio necesita los opcionales de esta
            # persona para poder mostrarlos/editarlos (antes se editaba inline, sin necesitarlos).
            "extra_fields": u.get_extras(),
            "categories": categories_by_user.get(u.id, []),
        })

    return result

@router.get("/email-check")
async def email_check(
    event_id: int, email: str, exclude_id: str = "", db: Session = Depends(get_db),
    staff: StaffUser = Depends(require_role("digitador")),
):
    """Aviso de correo ya existente (reunión 2026-09-21, ítem 20) — SOLO lectura: dice si otra
    persona de este cliente (las personas viven a nivel de tenant) ya tiene ese correo, para que
    quien digita lo vea ANTES de guardar. No bloquea nada. `exclude_id` = la propia persona al
    editar, para que su propio correo no cuente como duplicado."""
    event = get_event_for_staff(event_id, db, staff)
    wanted = email.strip().lower()
    if not wanted:
        return {"exists": False, "matches": []}
    query = db.query(User).filter(User.tenant_id == event.tenant_id, func.lower(User.email) == wanted)
    if exclude_id:
        query = query.filter(User.id != exclude_id)
    matches = [{"id": u.id, "name": f"{u.first_name or ''} {u.last_name or ''}".strip()} for u in query.limit(5).all()]
    return {"exists": bool(matches), "matches": matches}


@router.post("/recognize")
async def recognize(
    event_id: int = Form(...), file: UploadFile = File(...), force: bool = Form(False),
    confirm: bool = Form(False),
    db: Session = Depends(get_db), staff: StaffUser = Depends(require_role("digitador")),
):
    """Sprint 2.2 Fase B (2026-09-16): un match facial YA NO acredita solo por defecto — antes
    creaba el AccessLog apenas encontraba la cara, sin que el digitador confirmara nada. Ahora,
    salvo que `Event.auto_register` esté prendido (switch por evento) o venga `confirm=true` (el
    digitador ya confirmó en el modal "Guardar y autorizar acceso"), un match devuelve
    `result: "MATCH_PENDING"` con los datos de la persona SIN crear ningún log todavía — el
    frontend reenvía la MISMA petición (mismo `file`, vía FormData reusado) con `confirm=true`
    para recién ahí acreditar de verdad. El flujo DUPLICADO/force sigue exactamente igual, se
    evalúa ANTES de este chequeo nuevo."""
    event = get_event_for_staff(event_id, db, staff)
    require_event_in_progress(event)
    img_array = BiometricEngine.process_image_stream(await file.read())
    unknown_enc = BiometricEngine.extract_encoding(img_array)

    if not unknown_enc:
        return {"result": "NO", "details": "Rostro no detectado"}

    users = db.query(User).filter(User.tenant_id == event.tenant_id).all()
    for user in users:
        known_enc = user.get_encoding()
        if known_enc and BiometricEngine.compare(known_enc, unknown_enc):
            if not force and _already_checked_in(db, event.id, user.id):
                return _duplicate_warning(db, event.id, user)
            data = {
                "id": user.id, "first_name": user.first_name, "last_name": user.last_name,
                "role": user.role, "entity": user.entity, "phone": user.phone,
                "email": user.email, "opt_1": user.opt_1, "opt_2": user.opt_2
            }
            if not event.auto_register and not confirm:
                return {"result": "MATCH_PENDING", "data": data}
            log = AccessLog(
                tenant_id=event.tenant_id, user_id=user.id, record_type="Existente",
                event_id=event.id, registered_by_staff_id=staff.id,
                registration_method="biometrico",  # Fase 16: este endpoint es SIEMPRE reconocimiento facial
            )
            db.add(log)
            _upsert_attendee(db, event.id, user.id, event.tenant_id)
            db.commit()
            return {"result": "SÍ", "data": data}

    return {"result": "NO", "details": "Denegado"}

def _identity_name_matches(db_name: str, scanned_name: str) -> bool:
    """Fuzzy-match de identidad (2026-09-16, Sprint 2.4 Fase 3 — corrección real) — distinto de
    `_matches_by_word_prefix` (que exige que TODAS las palabras de lo que se busca prefijen algo
    en el texto largo, pensado para búsquedas tipo "escribo un pedazo del nombre de un evento").
    Acá el caso real es al revés y en ambas direcciones a la vez: la base puede tener el nombre
    abreviado y lo escaneado el nombre completo real (`JHOAN SEBASTIAN ANGARITA ROJAS` escaneado
    vs. `Sebas Angarita` en la base), o viceversa (`Sebas Angarita` escaneado vs. `Sebastián
    Angarita` en la base, el caso original). Por eso se exige que CADA palabra del nombre en BASE
    (el lado casi siempre más corto/abreviado) tenga alguna palabra en lo escaneado que la
    prefije O que sea prefijada por ella — no al revés, y no en una sola dirección."""
    db_words = _words(db_name)
    scanned_words = _words(scanned_name)
    if not db_words or not scanned_words:
        return False
    # Ítem 2 (reunión 2026-09-21): además del prefijo en ambas direcciones, tolera errores de
    # tipeo ("Yuliana" en base vs "Juliana" escaneado) — ver `_typo_match`.
    return all(
        any(sw.startswith(dw) or dw.startswith(sw) or _typo_match(dw, sw) for sw in scanned_words)
        for dw in db_words
    )


@router.post("/checkin-cedula")
async def checkin_cedula(
    event_id: int = Form(...), cedula: str = Form(...), force: bool = Form(False), confirm: bool = Form(False),
    first_name: str = Form(""), last_name: str = Form(""),
    db: Session = Depends(get_db), staff: StaffUser = Depends(require_role("digitador")),
):
    """Acreditación por cédula (lector de código de barras o MRZ de la cédula nueva) — mismo
    shape de respuesta que /recognize (result + data), para reusar el mismo patrón de frontend.
    Si la persona ya es conocida en el tenant (estuviera o no precargada para este evento
    puntual), se acredita directo como 'Existente' y de paso queda asociada a este evento. Si ya
    tiene un AccessLog para este evento, se avisa (result: DUPLICADO) en vez de acreditar de
    nuevo, salvo que venga force=true (el operador ya confirmó que sí quiere repetirlo).

    Sprint 2.4, Fase 2/3 (2026-09-16, pedido explícito, corregido en la Fase 3): un match EXACTO
    por cédula SIEMPRE deja la fila filtrada/visible en el Directorio en Vivo, sin modal — pero
    solo ACREDITA de una vez (crea el AccessLog, la fila sale "verde") si `Event.auto_register`
    está prendido. Si está apagado, se devuelve `result: "FOUND_PENDING"` — la persona queda
    visible pero SIN acreditar (blanco/"No registrado"), y el frontend ofrece un botón puntual
    "Acreditar" en esa fila para confirmar con un clic (reintenta esta misma petición con
    `confirm=true`) — sin volver a un modal flotante aparte. El reconocimiento facial
    (`/recognize`) sigue con su propio flujo de confirmación, sin cambios, no es parte de esto.

    Si NO hay match exacto por cédula, ya NO se ofrece alta manual automática — se intenta un
    fallback por nombre (`first_name`/`last_name`, cuando el frontend los tiene: vienen del CSV
    de la cédula vieja o del OCR de la MRZ de la nueva) contra TODOS los `User` de este tenant,
    con `_identity_name_matches` (ver arriba — bidireccional, no `_matches_by_word_prefix`) — se
    devuelven como candidatos (`name_matches`), NINGUNO se acredita solo, es el operador quien
    decide desde el Directorio filtrado. Sin nombre (cédula suelta sin match), no hay con qué
    intentar el fallback."""
    cedula = cedula.strip()
    event = get_event_for_staff(event_id, db, staff)
    require_event_in_progress(event)

    user = db.query(User).filter(User.id == cedula, User.tenant_id == event.tenant_id).first()
    if not user:
        full_name = f"{first_name} {last_name}".strip()
        name_matches = []
        if full_name:
            candidates = db.query(User).filter(User.tenant_id == event.tenant_id).all()
            for candidate in candidates:
                db_name = f"{candidate.first_name or ''} {candidate.last_name or ''}"
                if _identity_name_matches(db_name, full_name):
                    name_matches.append({"id": candidate.id, "first_name": candidate.first_name, "last_name": candidate.last_name})
        return {"result": "NO_MATCH", "details": "Cédula no encontrada", "name_matches": name_matches}

    if not force and _already_checked_in(db, event.id, user.id):
        return _duplicate_warning(db, event.id, user)

    data = {
        "id": user.id, "first_name": user.first_name, "last_name": user.last_name,
        "role": user.role, "entity": user.entity, "phone": user.phone,
        "email": user.email, "opt_1": user.opt_1, "opt_2": user.opt_2
    }

    if not event.auto_register and not confirm:
        return {"result": "FOUND_PENDING", "data": data}

    log = AccessLog(
        tenant_id=event.tenant_id, user_id=user.id, record_type="Existente",
        event_id=event.id, registered_by_staff_id=staff.id,
        # Fase 16 (2026-09-17): "tradicional" = cédula encontrada y confirmada a mano;
        # "autoregistro" = el mismo match, pero acreditado solo porque el evento tiene el switch
        # prendido — misma cédula, la diferencia es si hizo falta que alguien confirmara.
        registration_method="autoregistro" if event.auto_register else "tradicional",
    )
    db.add(log)
    _upsert_attendee(db, event.id, user.id, event.tenant_id)
    db.commit()
    return {"result": "SÍ", "data": data}

@router.patch("/users/{user_id}")
async def update_user(
    user_id: str, data: dict, event_id: int, db: Session = Depends(get_db),
    staff: StaffUser = Depends(require_role("coordinador")),
):
    event = get_event_for_staff(event_id, db, staff)
    user = db.query(User).filter(User.id == user_id, User.tenant_id == event.tenant_id).first()
    if user:
        # extra_fields llega como un dict {opcional_N: valor} desde el modal de "Editar" del
        # Directorio (2026-09-16) — no puede pasar por el setattr genérico de abajo: la columna
        # real (`User.extra_fields`) es un Text con JSON serializado a mano (`get_extras()`/
        # `set_extras()`), no un dict crudo; asignarlo directo lo corrompería.
        categories = data.pop("categories", None)  # ítem 14: por evento (EventAttendee), no una columna de User
        if categories is not None:
            _upsert_attendee(db, event.id, user.id, event.tenant_id, _clean_categories(event, categories))
        extra_fields = data.pop("extra_fields", None)
        if extra_fields is not None:
            merged = user.get_extras()
            merged.update(extra_fields)
            user.set_extras(merged)
        # Sprint 2.4 Fase 7 (2026-09-16, pedido explícito: "que no se pueda repetir para nadie el
        # id"): 'id'/'tenant_id' son la llave primaria compuesta de User, referenciada por FK
        # desde EventAttendee/AccessLog — un setattr genérico sobre ellas intentaría un UPDATE de
        # la propia PK, algo que este endpoint nunca tuvo pensado hacer (eso es exactamente lo que
        # resuelve PUT /users/{user_id}/cedula con su patrón seguro insertar-reapuntar-borrar, ver
        # abajo) y que solo terminaría en un IntegrityError sin manejar. Se ignoran acá — el único
        # camino válido para cambiar la cédula es ese otro endpoint.
        data.pop("id", None)
        data.pop("tenant_id", None)
        for key, value in data.items():
            if hasattr(user, key):
                setattr(user, key, value)

        field_configs = parametros.field_configs_for_event(db, event)
        final_values = {
            "role": user.role, "entity": user.entity, "phone": user.phone,
            "email": user.email, "opt_1": user.opt_1, **user.get_extras(),
        }
        missing = _missing_required_fields(field_configs, final_values)
        if missing:
            raise HTTPException(status_code=400, detail=f"Faltan campos obligatorios: {', '.join(missing)}")

        log = AccessLog(
            tenant_id=event.tenant_id, user_id=user.id, record_type="Actualizado",
            event_id=event.id, registered_by_staff_id=staff.id,
        )
        db.add(log)
        db.commit()
        return {"message": "Actualizado correctamente"}
    return {"error": "Usuario no encontrado"}

@router.put("/users/{user_id}/cedula")
async def update_user_cedula(
    user_id: str, event_id: int, data: dict, db: Session = Depends(get_db),
    staff: StaffUser = Depends(require_role("admin")),
):
    """Corrige la cédula (User.id) de una persona ya registrada (2026-09-16, pedido explícito: a
    veces se acredita por nombre porque la cédula quedó mal escrita, y no había forma de
    arreglarla). User.id es parte de la llave primaria compuesta (id, tenant_id) y está referenciada
    por FK desde EventAttendee/AccessLog — renombrarla in-place violaría esa FK a mitad de camino,
    así que el patrón seguro es insertar una fila nueva con el id correcto, reapuntar las filas
    hijas, y recién ahí borrar la vieja, todo en una sola transacción. admin+ solamente: User vive
    a nivel de tenant (se comparte entre eventos, ver CLAUDE.md), así que este cambio afecta a la
    persona en TODOS los eventos de este cliente, no solo en este — más sensible que un campo
    normal del perfil."""
    event = get_event_for_staff(event_id, db, staff)
    old_id = user_id
    new_id = str(data.get("new_id", "")).strip()
    if not new_id:
        raise HTTPException(status_code=400, detail="La nueva cédula no puede estar vacía")
    if new_id == old_id:
        raise HTTPException(status_code=400, detail="La nueva cédula es igual a la actual")

    user = db.query(User).filter(User.id == old_id, User.tenant_id == event.tenant_id).first()
    if not user:
        raise HTTPException(status_code=404, detail="Persona no encontrada")

    conflict = db.query(User).filter(User.id == new_id, User.tenant_id == event.tenant_id).first()
    if conflict:
        raise HTTPException(status_code=409, detail="Ya existe otra persona con esa cédula en este cliente")

    new_user = User(
        id=new_id, tenant_id=user.tenant_id, first_name=user.first_name, last_name=user.last_name,
        role=user.role, entity=user.entity, phone=user.phone, email=user.email,
        opt_1=user.opt_1, opt_2=user.opt_2, extra_fields=user.extra_fields, face_encoding=user.face_encoding,
    )
    db.add(new_user)
    db.flush()  # el nuevo User debe existir antes de reapuntar las filas hijas hacia él

    db.query(EventAttendee).filter(
        EventAttendee.user_id == old_id, EventAttendee.tenant_id == event.tenant_id
    ).update({"user_id": new_id}, synchronize_session=False)
    db.query(AccessLog).filter(
        AccessLog.user_id == old_id, AccessLog.tenant_id == event.tenant_id
    ).update({"user_id": new_id}, synchronize_session=False)
    # PrintLog (Sprint 2.4 Fase 3, 2026-09-16) tiene la misma FK compuesta (user_id, tenant_id)
    # que EventAttendee/AccessLog — bug real encontrado en Fase 7 (2026-09-16): faltaba
    # reapuntarlo acá también, así que renombrar la cédula de alguien que ya se había impreso
    # antes tumbaba este endpoint con un ForeignKeyViolation al borrar el User viejo más abajo.
    db.query(PrintLog).filter(
        PrintLog.user_id == old_id, PrintLog.tenant_id == event.tenant_id
    ).update({"user_id": new_id}, synchronize_session=False)
    db.flush()

    db.delete(user)

    signatures.rename_signatures(event.tenant_id, old_id, new_id)
    old_photo = os.path.join(_known_faces_dir(event.tenant_id), f"{old_id}.jpg")
    if os.path.exists(old_photo):
        os.rename(old_photo, os.path.join(_known_faces_dir(event.tenant_id), f"{new_id}.jpg"))

    db.commit()
    return {"message": "Cédula actualizada correctamente", "new_id": new_id}


@router.delete("/users/{user_id}/logs")
async def delete_user_logs(
    user_id: str, event_id: int, db: Session = Depends(get_db),
    staff: StaffUser = Depends(require_role("admin")),
):
    """Histórico (2026-09-21 en adelante) — ya NO se usa desde el frontend, quedó reemplazado por
    `DELETE /users/{id}` (borrado real, ver abajo) y `PATCH /events/{id}/users/{id}/status`
    (cambio de estado sin borrar nada). Se deja documentado, no se borra el endpoint, por si algún
    integrador externo llegó a depender de él. Ojo: este SÍ borra por tenant completo, no por
    evento — bug conocido, no replicar este patrón en código nuevo."""
    event = get_event_for_staff(event_id, db, staff)
    db.query(AccessLog).filter(AccessLog.user_id == user_id, AccessLog.tenant_id == event.tenant_id).delete()
    db.commit()
    return {"message": "Registros eliminados. Estado regresado a No Registrado."}


@router.delete("/users/{user_id}")
async def delete_user_from_event(
    user_id: str, event_id: int, db: Session = Depends(get_db),
    staff: StaffUser = Depends(require_role("coordinador")),
):
    """Borra a la persona de ESTE evento por completo (2026-09-16, pedido explícito: "Eliminar"
    debía borrar de verdad, no solo resetear el estado). Scoped al evento — no confundir con el
    viejo `.../logs` de arriba, que por bug afectaba TODO el tenant. Si tras sacarla de este
    evento la persona no queda en NINGÚN otro evento del mismo tenant (mismo User se reusa entre
    eventos, ver CLAUDE.md), se borra también el `User` y su foto física — si sigue en otro
    evento, ese `User` se conserva intacto, no se puede reventar la base de un evento ajeno."""
    event = get_event_for_staff(event_id, db, staff)
    user = db.query(User).filter(User.id == user_id, User.tenant_id == event.tenant_id).first()
    if not user:
        raise HTTPException(status_code=404, detail="Persona no encontrada")

    db.query(EventAttendee).filter(EventAttendee.event_id == event_id, EventAttendee.user_id == user_id).delete()
    db.query(AccessLog).filter(AccessLog.event_id == event_id, AccessLog.user_id == user_id).delete()
    signatures.delete_signatures(event.tenant_id, event_id, user_id)
    db.flush()

    other_attendee = db.query(EventAttendee).filter(
        EventAttendee.user_id == user_id, EventAttendee.tenant_id == event.tenant_id
    ).first()
    other_log = db.query(AccessLog).filter(
        AccessLog.user_id == user_id, AccessLog.tenant_id == event.tenant_id
    ).first()

    fully_deleted = False
    if not other_attendee and not other_log:
        photo_path = os.path.join(_known_faces_dir(event.tenant_id), f"{user_id}.jpg")
        if os.path.exists(photo_path):
            os.remove(photo_path)
        db.delete(user)
        fully_deleted = True

    db.commit()
    return {
        "message": "Persona eliminada de este evento" + (
            " y de la base de datos (no pertenecía a ningún otro evento de este cliente)."
            if fully_deleted else ". Sigue en la base porque pertenece a otro evento del mismo cliente."
        ),
        "fully_deleted": fully_deleted,
    }


@router.patch("/events/{event_id}/users/{user_id}/status")
async def update_registration_status(
    event_id: int, user_id: str, data: dict, db: Session = Depends(get_db),
    staff: StaffUser = Depends(require_role("coordinador")),
):
    """Cambia manualmente el estado de registro de una persona en este evento (2026-09-16, nuevo;
    ampliado a coordinador+ el 2026-09-17, pedido explícito — antes admin+ solamente, junto con
    DELETE /users/{id} arriba) — antes la única forma de pasar a "Registrado" era un escaneo/
    búsqueda real, y no existía forma de volver a "No registrado" salvo el viejo borrado-de-logs.
    Pensado para el caso real que describió Juan David: crear a alguien de antemano (aún no ha
    llegado) sin que quede "Registrado" de una, y poder marcarlo cuando sí llegue — o al revés,
    corregir un registro hecho por error. `data: {"status": "registrado" | "no_registrado"}`."""
    event = get_event_for_staff(event_id, db, staff)
    user = db.query(User).filter(User.id == user_id, User.tenant_id == event.tenant_id).first()
    if not user:
        raise HTTPException(status_code=404, detail="Persona no encontrada")

    new_status = data.get("status")
    if new_status not in ("registrado", "no_registrado"):
        raise HTTPException(status_code=400, detail="status debe ser 'registrado' o 'no_registrado'")

    if new_status == "registrado":
        if not _already_checked_in(db, event_id, user_id):
            log = AccessLog(
                tenant_id=event.tenant_id, user_id=user_id, record_type="Existente",
                event_id=event_id, registered_by_staff_id=staff.id,
                registration_method="tradicional",  # Fase 16: cambio manual de estado desde el Directorio
            )
            db.add(log)
            _upsert_attendee(db, event_id, user_id, event.tenant_id)
    else:
        db.query(AccessLog).filter(AccessLog.event_id == event_id, AccessLog.user_id == user_id).delete()

    db.commit()
    return {"message": "Estado de registro actualizado", "status": new_status}

@router.post("/register")
async def manual_register(
    event_id: int = Form(...), id: str = Form(...), first_name: str = Form(...), last_name: str = Form(...),
    role: str = Form(""), entity: str = Form(""), phone: str = Form(""),
    email: str = Form(""), opt_1: str = Form(""), extra_fields: str = Form(None), categories: str = Form(None),
    field_labels: str = Form(None), file: UploadFile = File(None), force: bool = Form(False),
    db: Session = Depends(get_db), staff: StaffUser = Depends(require_role("digitador")),
):
    """file es OPCIONAL: el alta manual la puede disparar tanto el flujo facial (con foto, para
    poder reconocer a esta persona después) como el de cédula (sin foto, solo identidad).
    Si la cédula ya existe como User en el tenant (de este evento o de uno anterior — User vive
    a nivel de tenant, no de evento) no se rechaza de plano: si todavía no tiene AccessLog para
    ESTE evento, simplemente se le agrega (reutilizar a alguien de un evento pasado es válido);
    si ya lo tiene, se avisa (DUPLICADO) igual que recognize/checkin-cedula, salvo force=true.

    Campos opcionales (2026-09-22, mismo esquema que bulk_register): `extra_fields` es un JSON
    {"opcional_1": "valor", ...} armado por el frontend (kiosk_registro.html deja agregar hasta
    30 campos opcionales al formulario de alta individual, con los mismos nombres que el roster
    original si lo hay). Si alguna clave usada todavía no tiene rótulo para este evento, se
    devuelve {"result": "NEEDS_LABELS", "fields": [...]} igual que bulk_register — el frontend
    normalmente ya manda `field_labels` de una vez porque él mismo preguntó el nombre al agregar
    el campo, así que este camino solo se ejerce en casos raros/defensivos."""
    event = get_event_for_staff(event_id, db, staff)
    require_event_in_progress(event)

    extras = _parse_extra_fields(extra_fields)
    try:
        person_categories = _clean_categories(event, json.loads(categories)) if categories else []
    except ValueError:
        person_categories = []
    used_optional_keys = set(extras.keys())
    missing = used_optional_keys - set(event.get_optional_labels().keys())
    if missing and not field_labels:
        return {"result": "NEEDS_LABELS", "fields": sorted(missing, key=lambda k: int(k.split("_")[1]))}
    if field_labels:
        _apply_optional_labels(event, used_optional_keys, field_labels)
        db.commit()

    existing = db.query(User).filter(User.id == id, User.tenant_id == event.tenant_id).first()
    if existing:
        if not force and _already_checked_in(db, event.id, existing.id):
            return _duplicate_warning(db, event.id, existing)
        log = AccessLog(
            tenant_id=event.tenant_id, user_id=existing.id, record_type="Existente",
            event_id=event.id, registered_by_staff_id=staff.id,
            registration_method="tradicional",  # Fase 16: alta manual, formulario
        )
        db.add(log)
        _upsert_attendee(db, event.id, existing.id, event.tenant_id, person_categories if categories else None)
        db.commit()
        return {"message": "Esta persona ya existía en el sistema — registrada para este evento."}

    field_configs = parametros.field_configs_for_event(db, event)
    final_values = {"role": role, "entity": entity, "phone": phone, "email": email, "opt_1": opt_1, **extras}
    missing = _missing_required_fields(field_configs, final_values)
    if missing:
        raise HTTPException(status_code=400, detail=f"Faltan campos obligatorios: {', '.join(missing)}")

    face_enc_json = None
    img_array = None
    # OJO: un <input type="file"> vacío igual manda una parte multipart (con filename=""), así
    # que `file is not None` es cierto incluso sin archivo real — hay que revisar `file.filename`
    # también, si no, `file.read()` da bytes vacíos y `Image.open()` truena (500) en vez de
    # tratarlo como "no se adjuntó foto" (CEDULA-08, bug real encontrado en testing 2026-09-21).
    if file is not None and file.filename:
        img_array = BiometricEngine.process_image_stream(await file.read())
        encodings = BiometricEngine.extract_encoding(img_array, is_registration=True)
        if not encodings:
            return {"error": "No se detectó un rostro en la fotografía."}
        face_enc_json = json.dumps(encodings)

    user = User(
        id=id, tenant_id=event.tenant_id, first_name=first_name.strip(), last_name=last_name.strip(),
        role=role, entity=entity, phone=phone, email=email, opt_1=opt_1, face_encoding=face_enc_json
    )
    user.set_extras(extras)
    db.add(user)
    # Mismo bug que ya se había arreglado en bulk_register (ver ese comentario) pero nunca se
    # replicó aquí: con autoflush=False, el INSERT de User quedaba pendiente sin mandarse a la
    # base todavía cuando AccessLog/EventAttendee (que dependen de él por llave foránea) se
    # intentaban insertar en el mismo flush — Postgres (y SQLite con PRAGMA foreign_keys=ON)
    # rechazan el INSERT de event_attendees con ForeignKeyViolation porque el User referenciado
    # técnicamente "no existe todavía" en ese punto. Un flush() explícito aquí garantiza que el
    # INSERT de User ya se mandó antes de crear las filas que dependen de él (reproducido y
    # confirmado con un test end-to-end vía FastAPI TestClient + SQLite con FK activas, 2026-09-23).
    db.flush()

    log = AccessLog(
        tenant_id=event.tenant_id, user_id=id, record_type="Nuevo",
        event_id=event.id, registered_by_staff_id=staff.id,
        registration_method="tradicional",  # Fase 16: alta manual, formulario
    )
    db.add(log)
    _upsert_attendee(db, event.id, id, event.tenant_id, person_categories)
    db.commit()

    if img_array is not None:
        img_path = os.path.join(_known_faces_dir(event.tenant_id), f"{id}.jpg")
        Image.fromarray(img_array).save(img_path)

    return {"message": "Usuario registrado exitosamente como Nuevo."}

@router.post("/bulk_register")
async def bulk_register(
    event_id: int = Form(...), roster_file: UploadFile = File(...), zip_file: UploadFile = File(None),
    field_labels: str = Form(None),
    db: Session = Depends(get_db), staff: StaffUser = Depends(require_role_excluding("coordinador", ("comercial",))),
):
    """Carga la base de asistentes esperados para el evento — sirve para CUALQUIER método de
    registro (cédula, facial, QR futuro), no es exclusiva de facial. roster_file acepta .csv o
    .xlsx (Historia 1.1: el negocio manda Excel). zip_file es OPCIONAL: solo hace falta si además
    quieres que estas personas se puedan reconocer por cara (las fotos del zip, nombradas
    <cédula>.jpg, se procesan a encoding biométrico). Sin zip, solo se cargan datos de identidad +
    se asocian al evento (EventAttendee) — suficiente para acreditar por cédula. Los errores de
    fila no abortan la carga completa, se reportan al final. No se permite cargar sobre un evento
    ya 'finalizado' (sí sobre 'creado' o 'en_proceso' — es preparación previa al evento).

    Campos dinámicos "opcional_1".."opcional_30" (2026-09-20): además de las columnas fijas
    (id, nombres, apellidos, cargo, entidad, telefono, correo, "tipo de asistente"), el roster
    puede traer hasta 30 columnas "opcional_N" — se interpretan las que de verdad vengan usadas
    (con al menos un valor no vacío en alguna fila), el resto se ignoran. Si el archivo trae
    columnas opcionales que este evento todavía no tiene rotuladas (Event.optional_field_labels),
    se devuelve {"result": "NEEDS_LABELS", "fields": [...]} SIN procesar nada — el frontend debe
    preguntarle al operador a qué corresponde cada una y reenviar la misma petición con
    field_labels (JSON, ej. {"opcional_1": "Talla de camisa"}) para que se guarden y se continúe
    con la carga."""
    event = get_event_for_staff(event_id, db, staff)
    if event.status == "finalizado":
        raise HTTPException(status_code=400, detail="No se puede cargar la base de un evento finalizado")
    # Bloqueo real (no solo el aviso que ya hace el frontend antes de mandar la petición, ver
    # kiosk_roster.html) — un evento en_proceso que YA tuvo una carga exitosa antes no puede
    # recibir otra: es casi siempre alguien resubiendo por error, y pisaría/duplicaría registros
    # de gente que ya se acreditó en vivo. Si de verdad necesitan cargar otra base, la salida es
    # crear un evento nuevo, no reintentar sobre este (2026-09-22, pedido explícito del usuario).
    if event.status == "en_proceso" and event.roster_uploaded:
        raise HTTPException(
            status_code=400,
            detail=(
                "Este evento ya tiene una base cargada y está EN PROCESO. Para evitar conflictos "
                "con los registros que ya se hicieron en vivo, no se puede volver a cargar otra "
                "base sobre este evento — si necesitas cargar una base distinta, crea un evento nuevo."
            ),
        )

    content = await roster_file.read()
    header_columns, rows = _read_roster_rows(roster_file.filename, content)

    # 2026-09-16, pedido explícito: antes, un archivo con columnas completamente distintas a la
    # plantilla se procesaba igual — cada fila terminaba reportada como "❌ sin ID/cédula" (porque
    # ninguna columna reconocida traía nada), pero la carga "completaba" con 0 perfiles sin dejar
    # claro que el problema real era el FORMATO del archivo, no los datos. Ahora se valida de una
    # que las columnas mínimas de la plantilla estén presentes antes de procesar ninguna fila.
    has_id_col = any(k in header_columns for k in _ID_KEYS)
    has_nombres_col = any(k in header_columns for k in ('nombres', 'nombre'))
    has_apellidos_col = any(k in header_columns for k in ('apellidos', 'apellido'))
    if not rows or not (has_id_col and has_nombres_col and has_apellidos_col):
        raise HTTPException(
            status_code=400,
            detail=(
                "Este archivo no tiene el formato esperado — deben venir al menos las columnas "
                "'id' (o 'cédula'), 'nombres' y 'apellidos'. Descarga y usa la plantilla oficial "
                "de carga de asistentes en vez de un archivo con otras columnas."
            ),
        )

    used_optional_keys = set()
    for row in rows:
        for raw_key, value in row.items():
            norm = _normalize_optional_key(raw_key)
            if norm and (value or "").strip():
                used_optional_keys.add(norm)

    missing = used_optional_keys - set(event.get_optional_labels().keys())
    if missing and not field_labels:
        # 2026-09-16, pedido explícito: hasta 2 valores de ejemplo por campo (tal como vienen en
        # el archivo) para que el operador identifique a qué corresponde sin tener que abrir el
        # Excel — ej. si "opcional_1" trae "Talla M"/"Talla L", eso mismo se muestra al preguntar.
        examples = _optional_field_examples(rows, missing)
        return {
            "result": "NEEDS_LABELS",
            "fields": sorted(missing, key=lambda k: int(k.split("_")[1])),
            "examples": examples,
        }

    if field_labels:
        _apply_optional_labels(event, used_optional_keys, field_labels)
        db.commit()

    known_faces_dir = _known_faces_dir(event.tenant_id)
    # Se declara acá (antes solo existía más abajo, en el pre-escaneo) para que el bloque del zip
    # de abajo pueda reportar en la misma lista — mismo criterio de prefijos (❌/⚠️/ℹ️) que el
    # resto del archivo, sin inventar un campo nuevo que el frontend no sepa mostrar.
    errors = []

    if zip_file is not None and zip_file.filename:
        zip_path = os.path.join(known_faces_dir, 'temp.zip')
        with open(zip_path, "wb") as buffer:
            buffer.write(await zip_file.read())

        with zipfile.ZipFile(zip_path, 'r') as zip_ref:
            for filename in zip_ref.namelist():
                if filename.startswith('__MACOSX') or filename.startswith('.') or filename.endswith('/'): continue
                basename = os.path.basename(filename)
                if not basename: continue
                with zip_ref.open(filename) as source:
                    try:
                        img_array = BiometricEngine.process_image_stream(source.read())
                    except Exception:
                        errors.append(f"⚠️ La foto '{basename}' del zip no se pudo leer (formato no soportado o archivo dañado) — no quedó asociada a nadie.")
                        continue
                    # Misma validación que ya hace /api/register (manual_register) para altas
                    # individuales: una foto sin rostro detectable no debe quedar guardada como si
                    # fuera una foto biométrica válida — antes se guardaba igual en silencio (bug
                    # real, QA local 2026-09-15: el coordinador nunca se enteraba de que esa
                    # persona quedó sin reconocimiento facial funcional).
                    encodings = BiometricEngine.extract_encoding(img_array, is_registration=True)
                    if not encodings:
                        errors.append(f"⚠️ La foto '{basename}' del zip no tiene un rostro detectable — no se asoció como foto biométrica de esa persona.")
                        continue
                    Image.fromarray(img_array).save(os.path.join(known_faces_dir, basename))
        if os.path.exists(zip_path): os.remove(zip_path)

        # Se enciende sola (nunca se apaga sola) — subir un roster sin zip más adelante no debe
        # quitarle a un evento la capacidad de reconocimiento facial que ya tenía. Decide si
        # /kiosk/{event_id}/registro muestra el escáner de cámara (2026-09-21, ver CLAUDE.md).
        if not event.facial_enabled:
            event.facial_enabled = True
            db.commit()

    # --- Pre-escaneo: identificar cédulas (con corrección del "Excel las volvió número"),
    # detectar cédulas repetidas dentro del mismo archivo, ANTES de tocar la base de datos.
    # Esto es lo que le permite al operador ver de una vez, con celda exacta, qué está mal en su
    # archivo en vez de enterarse fila por fila o con un error genérico. ---
    row_infos = []
    seen_rows_by_id = {}
    for row_num, clean_row in enumerate(rows, start=2):  # fila 1 es el encabezado
        identificador, id_cell, id_note = _extract_identificador(clean_row, header_columns, row_num)
        notes = [id_note] if id_note else []
        row_infos.append({
            "row_num": row_num, "clean_row": clean_row,
            "identificador": identificador, "id_cell": id_cell, "notes": notes,
        })
        if identificador:
            seen_rows_by_id.setdefault(identificador, []).append(row_num)

    id_col = next((header_columns[k] for k in _ID_KEYS if k in header_columns), None)
    for id_val, row_nums in seen_rows_by_id.items():
        if len(row_nums) > 1:
            cells = ", ".join(f"{id_col}{n}" if id_col else f"fila {n}" for n in row_nums)
            errors.append(
                f"⚠️ La cédula '{id_val}' aparece repetida en {len(row_nums)} filas ({cells}) — "
                f"se combinaron los datos y quedó lo de la última fila."
            )

    count = 0
    for info in row_infos:
        row_num, clean_row, identificador = info["row_num"], info["clean_row"], info["identificador"]
        errors.extend(info["notes"])

        if not identificador:
            where = info["id_cell"] or f"fila {row_num}"
            errors.append(f"❌ Fila {row_num} ({where}): sin ID/cédula, se omitió esta fila.")
            continue

        # Validaciones de calidad de dato — no bloquean la fila (se guarda igual), solo avisan
        # exactamente en qué celda está el problema para que el operador lo revise si quiere.
        nombres, nombres_cell = _field_lookup(clean_row, header_columns, row_num, 'nombres', 'nombre')
        if not nombres:
            errors.append(f"⚠️ Fila {row_num} ({nombres_cell or 'sin columna de nombres'}): falta el nombre.")
        apellidos, apellidos_cell = _field_lookup(clean_row, header_columns, row_num, 'apellidos', 'apellido')
        if not apellidos:
            errors.append(f"⚠️ Fila {row_num} ({apellidos_cell or 'sin columna de apellidos'}): falta el apellido.")
        correo, correo_cell = _field_lookup(clean_row, header_columns, row_num, 'correo', 'e-mail corporativo')
        if correo and not _EMAIL_RE.match(correo):
            errors.append(f"⚠️ Fila {row_num} ({correo_cell}): el correo '{correo}' no parece válido — se guardó igual, revísalo.")
        telefono, telefono_cell = _field_lookup(clean_row, header_columns, row_num, 'telefono', 'tel. celular')
        if telefono and not _PHONE_RE.match(telefono):
            errors.append(f"⚠️ Fila {row_num} ({telefono_cell}): el teléfono '{telefono}' tiene caracteres raros — se guardó igual, revísalo.")

        try:
            face_enc_json = None
            img_path = os.path.join(known_faces_dir, f"{identificador}.jpg")
            if os.path.exists(img_path):
                known_image = face_recognition.load_image_file(img_path)
                encodings = face_recognition.face_encodings(known_image, num_jitters=25)
                if encodings:
                    face_enc_json = json.dumps(encodings[0].tolist())

            # SAVEPOINT por fila (no solo un try/except de Python): con autoflush=False, una sola
            # fila con un problema real de base de datos (ej. una violación de llave) dejaba la
            # transacción completa "abortada" para Postgres, y el error solo aparecía hasta el
            # db.commit() final — tumbando TODA la carga con un 500, incluidas las filas que sí
            # habían "funcionado" en Python pero nunca llegaron a guardarse. Con begin_nested(),
            # si esta fila falla se revierte solo su savepoint (la transacción de afuera sigue
            # sana) y el flush() intermedio asegura que el INSERT de User ya se mandó a Postgres
            # antes del de EventAttendee (que depende de él por llave foránea).
            with db.begin_nested():
                user = db.query(User).filter(User.id == identificador, User.tenant_id == event.tenant_id).first()
                if not user:
                    user = User(id=identificador, tenant_id=event.tenant_id)
                    db.add(user)

                user.first_name = nombres
                user.last_name = apellidos
                user.role = clean_row.get('cargo', '')
                user.entity = clean_row.get('entidad', '') or clean_row.get('empresa', '')  # 'empresa' = columna del nombre viejo, sigue aceptándose en Excels ya armados
                user.phone = telefono
                user.email = correo
                user.opt_1 = clean_row.get('tipo de asistente', '') or clean_row.get('tipo_asistente', '')

                extras = {}
                for raw_key, value in clean_row.items():
                    norm = _normalize_optional_key(raw_key)
                    val = (value or "").strip()
                    if norm and val:
                        extras[norm] = val
                user.set_extras(extras)

                if face_enc_json:
                    user.face_encoding = face_enc_json

                db.flush()
                # Categorías (ítem 14): columna "categoria"/"categorias" (con o sin tilde), varias
                # separadas por coma/;/|. Las que el evento no conocía se agregan solas.
                raw_categories = next((clean_row[k] for k in ("categoria", "categorias", "categoría", "categorías") if clean_row.get(k)), None)
                _upsert_attendee(
                    db, event.id, identificador, event.tenant_id,
                    _clean_categories(event, raw_categories, auto_add=True) if raw_categories else None,
                )
                db.flush()

            count += 1
        except Exception as e:
            where = info["id_cell"] or f"fila {row_num}"
            errors.append(f"❌ Fila {row_num} ({where}): no se pudo guardar — {e}")

    if count > 0 and not event.roster_uploaded:
        event.roster_uploaded = True

    db.commit()
    message = f"Carga completa: {count} perfiles cargados."
    if errors:
        message += f" {len(errors)} observación(es) — revisa el detalle."
    return {"message": message, "count": count, "errors": errors, "optional_labels": event.get_optional_labels()}

@router.get("/report")
async def download_report(
    event_id: int, db: Session = Depends(get_db), staff: StaffUser = Depends(require_role("coordinador"))
    # 2026-09-17 (pedido explícito): 'cliente' ve Estadísticas (gráficos, vía
    # require_role_or_client en stats.py) pero YA NO puede exportar la base — antes usaba el mismo
    # require_role_or_client que las gráficas, ahora exige coordinador+ como cualquier otra
    # acción operativa, sin la excepción de cliente.
):
    event = get_event_for_staff(event_id, db, staff)
    temp_file = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
    ReportManager.generate_excel_report(db, event.id, event.tenant_id, temp_file.name)
    return FileResponse(temp_file.name, filename="Golden_Reporte_Eventos.xlsx")
