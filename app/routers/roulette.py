"""Ruleta / sorteos por evento (Sprint 5).

Dos configuraciones separadas: COMPORTAMIENTO (de dónde salen los candidatos y cómo se elige al ganador) y VISUAL
(cómo se ve la ruleta en la pantalla que ve la gente). La ejecución la dispara un coordinador+ desde su pantalla; la
pantalla de visualización (`/r/<token>`, sin login, para un proyector) consulta por sondeo y anima la ruleta
cuando aparece un sorteo nuevo. El resultado que se revela SIEMPRE es el ya determinado por el modo elegido:
predeterminado a mano en los modos 1-3 y 5a, realmente aleatorio (servidor, `SystemRandom`) solo en 5b y en "todos".

Modos (candidatos = siempre la base ya cargada del evento):
  single_fixed   1. un ganador elegido de antemano
  ordered        2. N ganadores en el orden exacto elegido
  shuffled       3. los mismos N elegidos, en orden aleatorio en cada ejecución
  all            4. todos los asistentes: cada ejecución saca a `batch_size` personas al azar entre quienes aún no
                 han salido en esta ronda (sin repetir) hasta agotarlos; "Reiniciar ronda" empieza de nuevo
  filter_manual  5a. el filtro solo acota la lista visible; el operador elige a mano (como 1-3)
  filter_random  5b. sorteo real: el sistema elige `count` al azar entre quienes cumplen el filtro

Auditoría (decisión: no estaba especificada, pero sin ella el módulo pierde valor): cada ejecución guarda un
`RouletteDraw` con nombre del sorteo, modo, si fue aleatorio real, filtro, ganadores, quién lo ejecutó y cuándo;
se consulta en pantalla y se descarga en Excel.
"""
import json
import os
import re
import secrets
import tempfile
import unicodedata
from datetime import datetime, timezone
from typing import Optional

from fastapi import APIRouter, Depends, File, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse, HTMLResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy.orm import Session

from app.auth import get_event_for_staff, require_role
from app.database import get_db
from app.models import AccessLog, Event, EventAttendee, RouletteConfig, RouletteDraw, StaffUser, User
from app.routers.badges import ALLOWED_IMAGE_EXT, _badge_assets_dir
from app.timeutil import to_local, _TZ

router = APIRouter()         # bajo /api (coordinador+)
public_router = APIRouter()  # /r/<token> (sin login)

MODES = ("single_fixed", "ordered", "shuffled", "all", "filter_manual", "filter_random")
_RNG = secrets.SystemRandom()
POOL_SIZE = 60  # nombres que se mandan a la pantalla para llenar la ruleta

ALLOWED_FONTS = [
    "Roboto", "Open Sans", "Lato", "Montserrat", "Oswald", "Raleway", "Poppins", "Playfair Display", "Inter",
    "Bebas Neue", "Anton", "Archivo Black", "Pacifico", "Lobster", "Dancing Script", "Merriweather", "Nunito",
    "Ubuntu", "Rubik", "Comfortaa", "Space Grotesk", "Manrope", "Zilla Slab",
]
DEFAULT_STYLE = {
    "title": "¡Sorteo!", "title_font": "Bebas Neue", "title_color": "#FFFFFF",
    "font_family": "Montserrat", "text_color": "#FFFFFF",
    "colors": ["#D4AF37", "#0A0E2E", "#B8941F", "#1B2456", "#E8C766", "#141B47"],
    "background_color": "#0A0E2E", "background_image": "", "logo": "", "pointer_color": "#FF4D6D",
    "winner_bg": "#D4AF37", "winner_color": "#0A0E2E", "spin_seconds": 6,
}
DEFAULT_BEHAVIOR = {"mode": "single_fixed", "winners": [], "count": 1, "batch_size": 1,
                    "filter": {"conditions": []}, "exclude_previous_winners": True, "all_reset_at": None}

_HEX = re.compile(r"^#[0-9a-fA-F]{6}$")


# ------------------------------------------------------------------ utilidades
def _norm(text) -> str:
    return "".join(c for c in unicodedata.normalize("NFD", str(text or "").lower()) if unicodedata.category(c) != "Mn").strip()


def _full_name(u) -> str:
    return f"{u.first_name or ''} {u.last_name or ''}".strip()


def _get_config(db: Session, event: Event, create=True) -> Optional[RouletteConfig]:
    cfg = db.query(RouletteConfig).filter_by(event_id=event.id).first()
    if not cfg and create:
        cfg = RouletteConfig(event_id=event.id, behavior_json=json.dumps(DEFAULT_BEHAVIOR), style_json=json.dumps(DEFAULT_STYLE))
        db.add(cfg)
        db.commit()
    return cfg


def _behavior(cfg) -> dict:
    return {**DEFAULT_BEHAVIOR, **(json.loads(cfg.behavior_json) if cfg.behavior_json else {})}


def _style(cfg) -> dict:
    return {**DEFAULT_STYLE, **(json.loads(cfg.style_json) if cfg.style_json else {})}


def candidates(db: Session, event: Event) -> list:
    """La base del evento (EventAttendee ∪ quienes ya tienen un registro real), con lo necesario para filtrar."""
    att = {a.user_id: a for a in db.query(EventAttendee).filter(EventAttendee.event_id == event.id)}
    logs = db.query(AccessLog).filter(AccessLog.event_id == event.id, AccessLog.record_type != "Actualizado").order_by(AccessLog.timestamp).all()
    first_log = {}
    for lg in logs:
        first_log.setdefault(lg.user_id, lg.timestamp)
    ids = set(att) | set(first_log)
    if not ids:
        return []
    out = []
    for u in db.query(User).filter(User.tenant_id == event.tenant_id, User.id.in_(ids)).order_by(User.last_name, User.first_name):
        a = att.get(u.id)
        out.append({
            "id": u.id, "name": _full_name(u), "entity": u.entity or "", "role": u.role or "", "opt_1": u.opt_1 or "",
            "extras": u.get_extras(), "categories": a.get_categories() if a else [],
            "registered_at": first_log.get(u.id), "status": "registrado" if u.id in first_log else "no_registrado",
        })
    return out


def _local_to_utc(value: str) -> datetime:
    """'2026-10-01T14:30' (hora local del evento) -> datetime naive en UTC."""
    local = datetime.fromisoformat(value).replace(tzinfo=_TZ)
    return local.astimezone(timezone.utc).replace(tzinfo=None)


def _match(person: dict, cond: dict) -> bool:
    field, op, value = cond.get("field"), cond.get("op", "equals"), cond.get("value")
    if field == "status":
        return person["status"] == value
    if field in ("registered_after", "registered_before"):
        when = person["registered_at"]
        if not when or not value:
            return False
        limit = _local_to_utc(value)
        return when >= limit if field == "registered_after" else when <= limit
    if field == "category":
        return any(_norm(c) == _norm(value) for c in person["categories"])
    if field in ("entity", "role", "opt_1"):
        target = person[field]
    elif isinstance(field, str) and field.startswith("extra:"):
        target = person["extras"].get(field[6:], "")
    else:
        raise HTTPException(status_code=400, detail=f"Campo de filtro no válido: {field}")
    a, b = _norm(target), _norm(value)
    return (b in a) if op == "contains" else (a == b)


def apply_filter(people: list, filt: Optional[dict]) -> list:
    conds = (filt or {}).get("conditions") or []
    return [p for p in people if all(_match(p, c) for c in conds)]


def _previous_winner_ids(db: Session, event: Event, since: Optional[datetime] = None, mode: Optional[str] = None) -> set:
    q = db.query(RouletteDraw).filter(RouletteDraw.event_id == event.id)
    if mode:
        q = q.filter(RouletteDraw.mode == mode)
    if since:
        q = q.filter(RouletteDraw.created_at > since)
    ids = set()
    for d in q:
        ids |= {w["id"] for w in json.loads(d.winners_json)}
    return ids


def _person_out(p: dict, position: int) -> dict:
    return {"position": position, "id": p["id"], "name": p["name"], "entity": p["entity"]}


def _serialize_draw(d: RouletteDraw, staff_names: dict) -> dict:
    return {
        "id": d.id, "label": d.label, "mode": d.mode, "is_random": d.is_random, "candidates_count": d.candidates_count,
        "filter": json.loads(d.filter_json) if d.filter_json else None, "winners": json.loads(d.winners_json),
        "created_at": to_local(d.created_at).strftime("%Y-%m-%dT%H:%M:%S"), "created_by": staff_names.get(d.created_by_id, ""),
    }


def _staff_names(db: Session, ids) -> dict:
    return {s.id: (s.full_name or s.username) for s in db.query(StaffUser).filter(StaffUser.id.in_(set(i for i in ids if i)))}


# ------------------------------------------------------------------ configuración
@router.get("/events/{event_id}/roulette")
async def get_roulette(event_id: int, db: Session = Depends(get_db), staff: StaffUser = Depends(require_role("coordinador"))):
    event = get_event_for_staff(event_id, db, staff)
    cfg = _get_config(db, event)
    people = candidates(db, event)
    return {
        "behavior": _behavior(cfg), "style": _style(cfg), "fonts": ALLOWED_FONTS, "modes": MODES,
        "has_display_link": bool(cfg.display_token), "candidates_total": len(people), "event_status": event.status,
    }


@router.get("/events/{event_id}/roulette/fields")
async def filter_fields(event_id: int, db: Session = Depends(get_db), staff: StaffUser = Depends(require_role("coordinador"))):
    """Campos por los que se puede filtrar: los fijos + cada campo opcional rotulado de este evento."""
    event = get_event_for_staff(event_id, db, staff)
    fields = [
        {"key": "status", "label": "Estado de registro", "type": "choice", "choices": [["registrado", "Registrado"], ["no_registrado", "No registrado"]]},
        {"key": "registered_after", "label": "Se registró desde (fecha y hora)", "type": "datetime"},
        {"key": "registered_before", "label": "Se registró hasta (fecha y hora)", "type": "datetime"},
        {"key": "category", "label": "Categoría", "type": "choice", "choices": [[c, c] for c in event.get_categories()]},
        {"key": "entity", "label": "Entidad", "type": "text"}, {"key": "role", "label": "Cargo", "type": "text"},
        {"key": "opt_1", "label": "Tipo de asistente", "type": "text"},
    ]
    for key, label in sorted(event.get_optional_labels().items(), key=lambda kv: int(kv[0].split("_")[1])):
        fields.append({"key": f"extra:{key}", "label": label, "type": "text"})
    return fields


@router.get("/events/{event_id}/roulette/candidates")
async def list_candidates(
    event_id: int, filter: Optional[str] = None, db: Session = Depends(get_db), staff: StaffUser = Depends(require_role("coordinador")),
):
    """Candidatos visibles (con el filtro opcional, JSON `{"conditions":[...]}`) — para elegir a mano (modos 1-3 y
    5a) y para ver cuántos cumplen el filtro (5b)."""
    event = get_event_for_staff(event_id, db, staff)
    people = candidates(db, event)
    if filter:
        try:
            people = apply_filter(people, json.loads(filter))
        except ValueError:
            raise HTTPException(status_code=400, detail="Filtro con formato inválido")
    return {"total": len(people), "people": [
        {"id": p["id"], "name": p["name"], "entity": p["entity"], "status": p["status"], "categories": p["categories"]} for p in people[:2000]
    ]}


def _validate_behavior(db: Session, event: Event, data: dict) -> dict:
    mode = data.get("mode")
    if mode not in MODES:
        raise HTTPException(status_code=400, detail=f"Modo inválido (uno de: {', '.join(MODES)})")
    people = candidates(db, event)
    by_id = {p["id"]: p for p in people}
    filt = data.get("filter") or {"conditions": []}
    winners = [str(w) for w in (data.get("winners") or [])]
    if len(set(winners)) != len(winners):
        raise HTTPException(status_code=400, detail="Un mismo ganador no puede repetirse en la lista")
    unknown = [w for w in winners if w not in by_id]
    if unknown:
        raise HTTPException(status_code=400, detail=f"Estas cédulas no están en la base del evento: {', '.join(unknown)}")
    if mode in ("single_fixed", "ordered", "shuffled", "filter_manual") and not winners:
        raise HTTPException(status_code=400, detail="Elige al menos un ganador")
    if mode == "single_fixed" and len(winners) != 1:
        raise HTTPException(status_code=400, detail="Este modo tiene exactamente un ganador")
    if mode in ("ordered", "shuffled") and len(winners) < 2:
        raise HTTPException(status_code=400, detail="Elige al menos 2 ganadores (para uno solo usa «Un ganador»)")
    if mode in ("filter_manual", "filter_random") and not (filt.get("conditions") or []):
        raise HTTPException(status_code=400, detail="Define al menos una condición de filtro")
    if mode == "filter_manual":
        allowed = {p["id"] for p in apply_filter(people, filt)}
        outside = [w for w in winners if w not in allowed]
        if outside:
            raise HTTPException(status_code=400, detail="Estos ganadores no cumplen el filtro: " + ", ".join(outside))
    try:
        count = int(data.get("count") or 1)
        batch = int(data.get("batch_size") or 1)
    except (TypeError, ValueError):
        raise HTTPException(status_code=400, detail="La cantidad debe ser un número")
    if count < 1 or batch < 1 or count > 500 or batch > 500:
        raise HTTPException(status_code=400, detail="La cantidad debe estar entre 1 y 500")
    return {"mode": mode, "winners": winners, "count": count, "batch_size": batch, "filter": {"conditions": filt.get("conditions") or []},
            "exclude_previous_winners": bool(data.get("exclude_previous_winners", True)), "all_reset_at": None}


@router.put("/events/{event_id}/roulette/behavior")
async def save_behavior(event_id: int, data: dict, db: Session = Depends(get_db), staff: StaffUser = Depends(require_role("coordinador"))):
    event = get_event_for_staff(event_id, db, staff)
    cfg = _get_config(db, event)
    new = _validate_behavior(db, event, data)
    old = _behavior(cfg)
    new["all_reset_at"] = old.get("all_reset_at")   # cambiar la configuración no reinicia la ronda de "todos"
    cfg.behavior_json = json.dumps(new)
    db.commit()
    return new


@router.put("/events/{event_id}/roulette/style")
async def save_style(event_id: int, data: dict, db: Session = Depends(get_db), staff: StaffUser = Depends(require_role("coordinador"))):
    event = get_event_for_staff(event_id, db, staff)
    cfg = _get_config(db, event)
    style = {**_style(cfg)}
    for key in ("title_color", "text_color", "background_color", "pointer_color", "winner_bg", "winner_color"):
        if key in data:
            if not _HEX.match(str(data[key])):
                raise HTTPException(status_code=400, detail=f"«{key}» debe ser un color como #D4AF37")
            style[key] = data[key]
    for key in ("title_font", "font_family"):
        if key in data:
            if data[key] not in ALLOWED_FONTS:
                raise HTTPException(status_code=400, detail="Fuente no disponible")
            style[key] = data[key]
    if "title" in data:
        style["title"] = str(data["title"])[:80]
    if "colors" in data:
        colors = data["colors"]
        if not isinstance(colors, list) or not (2 <= len(colors) <= 12) or not all(_HEX.match(str(c)) for c in colors):
            raise HTTPException(status_code=400, detail="Elige entre 2 y 12 colores válidos para los segmentos")
        style["colors"] = colors
    if "spin_seconds" in data:
        try:
            style["spin_seconds"] = max(2, min(20, int(data["spin_seconds"])))
        except (TypeError, ValueError):
            raise HTTPException(status_code=400, detail="La duración debe ser un número de segundos")
    for key in ("background_image", "logo"):
        if key in data:
            value = str(data[key] or "")
            if value and not value.startswith(event.tenant_id + "/"):
                raise HTTPException(status_code=400, detail="Imagen no válida")
            style[key] = value
    cfg.style_json = json.dumps(style)
    db.commit()
    return style


@router.post("/events/{event_id}/roulette/upload-image")
async def upload_image(event_id: int, file: UploadFile = File(...), db: Session = Depends(get_db), staff: StaffUser = Depends(require_role("coordinador"))):
    event = get_event_for_staff(event_id, db, staff)
    ext = os.path.splitext(file.filename or "")[1].lower()
    if ext not in ALLOWED_IMAGE_EXT:
        raise HTTPException(status_code=400, detail="Usa una imagen PNG, JPG, WEBP o GIF")
    content = await file.read()
    if len(content) > 8_000_000:
        raise HTTPException(status_code=400, detail="La imagen pesa más de 8 MB")
    name = f"{secrets.token_hex(16)}{ext}"
    with open(os.path.join(_badge_assets_dir(event.tenant_id), name), "wb") as f:
        f.write(content)
    return {"storage_path": f"{event.tenant_id}/{name}"}


@router.post("/events/{event_id}/roulette/display-link")
async def display_link(event_id: int, request: Request, regenerate: bool = False, db: Session = Depends(get_db), staff: StaffUser = Depends(require_role("coordinador"))):
    """Enlace de la pantalla de visualización (para el proyector): sin login, con código secreto. Se crea la primera
    vez; con `regenerate=true` el anterior deja de funcionar."""
    event = get_event_for_staff(event_id, db, staff)
    cfg = _get_config(db, event)
    if regenerate or not cfg.display_token:
        cfg.display_token = secrets.token_urlsafe(18)
        db.commit()
    base = (os.getenv("PUBLIC_BASE_URL") or str(request.base_url)).rstrip("/")
    return {"url": f"{base}/r/{cfg.display_token}"}


@router.get("/events/{event_id}/roulette/display-link")
async def get_display_link(event_id: int, request: Request, db: Session = Depends(get_db), staff: StaffUser = Depends(require_role("coordinador"))):
    event = get_event_for_staff(event_id, db, staff)
    cfg = _get_config(db, event)
    base = (os.getenv("PUBLIC_BASE_URL") or str(request.base_url)).rstrip("/")
    return {"url": f"{base}/r/{cfg.display_token}" if cfg.display_token else None}


# ------------------------------------------------------------------ ejecución
@router.post("/events/{event_id}/roulette/execute")
async def execute(event_id: int, data: dict, db: Session = Depends(get_db), staff: StaffUser = Depends(require_role("coordinador"))):
    """Ejecuta el sorteo con la configuración de comportamiento guardada. Devuelve los ganadores y guarda el registro;
    la pantalla de visualización lo recoge por sondeo y anima la ruleta."""
    event = get_event_for_staff(event_id, db, staff)
    cfg = _get_config(db, event)
    b = _behavior(cfg)
    mode = b["mode"]
    people = candidates(db, event)
    if not people:
        raise HTTPException(status_code=400, detail="Este evento todavía no tiene personas cargadas")
    by_id = {p["id"]: p for p in people}
    label = str(data.get("label") or "").strip()[:120]
    is_random, pool_people = False, people

    if mode in ("single_fixed", "ordered", "shuffled", "filter_manual"):
        chosen = [by_id[w] for w in b["winners"] if w in by_id]
        if not chosen:
            raise HTTPException(status_code=400, detail="Los ganadores elegidos ya no están en la base del evento — vuelve a elegirlos")
        if mode == "shuffled":
            _RNG.shuffle(chosen)
        if mode == "filter_manual":
            pool_people = apply_filter(people, b["filter"])
    elif mode == "all":
        since = datetime.fromisoformat(b["all_reset_at"]) if b.get("all_reset_at") else None
        drawn = _previous_winner_ids(db, event, since=since, mode="all")
        pool = [p for p in people if p["id"] not in drawn]
        if not pool:
            raise HTTPException(status_code=400, detail="Ya salieron todos los asistentes en esta ronda. Usa «Reiniciar ronda» para empezar de nuevo.")
        chosen = _RNG.sample(pool, min(b["batch_size"], len(pool)))
        is_random = True
    else:  # filter_random
        pool = apply_filter(people, b["filter"])
        pool_people = pool
        if b["exclude_previous_winners"]:
            gone = _previous_winner_ids(db, event)
            pool = [p for p in pool if p["id"] not in gone]
        if not pool:
            raise HTTPException(status_code=400, detail="Nadie cumple el filtro (o todos los que lo cumplen ya ganaron antes)")
        chosen = _RNG.sample(pool, min(b["count"], len(pool)))
        is_random = True

    winners = [_person_out(p, i) for i, p in enumerate(chosen, start=1)]
    draw = RouletteDraw(
        event_id=event.id, label=label, mode=mode, is_random=is_random, candidates_count=len(pool_people),
        filter_json=json.dumps(b["filter"]) if mode in ("filter_manual", "filter_random") else None,
        winners_json=json.dumps(winners), created_by_id=staff.id,
    )
    db.add(draw)
    db.commit()
    names = [p["name"] for p in pool_people]
    _RNG.shuffle(names)
    return {"draw_id": draw.id, "mode": mode, "is_random": is_random, "winners": winners, "pool": names[:POOL_SIZE], "label": label}


@router.post("/events/{event_id}/roulette/reset-round")
async def reset_round(event_id: int, db: Session = Depends(get_db), staff: StaffUser = Depends(require_role("coordinador"))):
    """Modo «todos»: empieza una ronda nueva (vuelven a poder salir todos)."""
    event = get_event_for_staff(event_id, db, staff)
    cfg = _get_config(db, event)
    b = _behavior(cfg)
    b["all_reset_at"] = datetime.utcnow().isoformat()
    cfg.behavior_json = json.dumps(b)
    db.commit()
    return {"message": "Ronda reiniciada"}


@router.get("/events/{event_id}/roulette/round")
async def round_status(event_id: int, db: Session = Depends(get_db), staff: StaffUser = Depends(require_role("coordinador"))):
    event = get_event_for_staff(event_id, db, staff)
    b = _behavior(_get_config(db, event))
    since = datetime.fromisoformat(b["all_reset_at"]) if b.get("all_reset_at") else None
    drawn = _previous_winner_ids(db, event, since=since, mode="all")
    return {"drawn": len(drawn), "total": len(candidates(db, event))}


# ------------------------------------------------------------------ historial y reporte
@router.get("/events/{event_id}/roulette/draws")
async def list_draws(event_id: int, db: Session = Depends(get_db), staff: StaffUser = Depends(require_role("coordinador"))):
    event = get_event_for_staff(event_id, db, staff)
    draws = db.query(RouletteDraw).filter_by(event_id=event.id).order_by(RouletteDraw.id.desc()).all()
    names = _staff_names(db, [d.created_by_id for d in draws])
    return [_serialize_draw(d, names) for d in draws]


MODE_LABELS = {
    "single_fixed": "Un ganador (elegido)", "ordered": "Varios, en orden", "shuffled": "Varios, orden aleatorio",
    "all": "Todos los asistentes", "filter_manual": "Filtro + elegido a mano", "filter_random": "Filtro + sorteo real",
}


@router.get("/events/{event_id}/roulette/report")
async def report(event_id: int, db: Session = Depends(get_db), staff: StaffUser = Depends(require_role("coordinador"))):
    """Reporte Excel de todos los sorteos del evento (mismo estilo que el resto de reportes)."""
    event = get_event_for_staff(event_id, db, staff)
    draws = db.query(RouletteDraw).filter_by(event_id=event.id).order_by(RouletteDraw.id).all()
    names = _staff_names(db, [d.created_by_id for d in draws])
    wb = Workbook()
    ws = wb.active
    ws.title = "Sorteos"
    ws.append([f"Sorteos — {event.name} ({event.event_code})"])
    ws["A1"].font = Font(bold=True, size=13, color="0A0E2E")
    headers = ["Fecha", "Hora", "Sorteo", "Modo", "¿Aleatorio real?", "Posición", "Cédula", "Nombre", "Entidad", "Candidatos", "Filtro", "Ejecutó"]
    ws.append(headers)
    for c in ws[2]:
        c.font = Font(color="FFFFFF", bold=True)
        c.fill = PatternFill(start_color="0A0E2E", end_color="0A0E2E", fill_type="solid")
        c.alignment = Alignment(horizontal="center")
    for d in draws:
        local = to_local(d.created_at)
        filt = "; ".join(f"{c['field']} {c.get('op', 'equals')} {c.get('value')}" for c in (json.loads(d.filter_json)["conditions"] if d.filter_json else []))
        for w in json.loads(d.winners_json):
            ws.append([local.strftime("%Y-%m-%d"), local.strftime("%H:%M:%S"), d.label or "", MODE_LABELS.get(d.mode, d.mode),
                       "Sí" if d.is_random else "No (elegido de antemano)", w["position"], w["id"], (w["name"] or "").upper(),
                       (w["entity"] or "").upper(), d.candidates_count, filt, names.get(d.created_by_id, "")])
    ws.auto_filter.ref = f"A2:{chr(64 + len(headers))}{ws.max_row}"
    for col, width in zip("ABCDEFGHIJKL", (12, 10, 28, 26, 24, 9, 14, 30, 24, 11, 34, 22)):
        ws.column_dimensions[col].width = width
    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
    wb.save(tmp.name)
    return FileResponse(tmp.name, filename=f"Sorteos_{event.event_code}.xlsx")


# ------------------------------------------------------------------ pantalla pública (proyector)
def _cfg_by_token(db: Session, token: str) -> RouletteConfig:
    cfg = db.query(RouletteConfig).filter(RouletteConfig.display_token == token).first()
    if not cfg:
        raise HTTPException(status_code=404, detail="Pantalla no disponible")
    return cfg


@public_router.get("/r/{token}", response_class=HTMLResponse)
async def display_page(token: str, request: Request, db: Session = Depends(get_db)):
    from app.main import templates  # import tardío: main.py importa este módulo
    try:
        cfg = _cfg_by_token(db, token)
    except HTTPException:
        return HTMLResponse("<h3 style='font-family:sans-serif;text-align:center;margin-top:20vh'>Pantalla no disponible</h3>", status_code=404)
    event = db.query(Event).filter(Event.id == cfg.event_id).first()
    return templates.TemplateResponse(request=request, name="roulette_display.html", context={"token": token, "event_name": event.name})


@public_router.get("/r/{token}/state")
async def display_state(token: str, after: Optional[int] = None, db: Session = Depends(get_db)):
    """Estilo + el `id` del último sorteo. La pantalla recuerda el `id` que ya vio y pide con `?after=<id>`: solo si hay
    uno MÁS NUEVO se arma y devuelve el sorteo (con la lista de nombres), así el sondeo de cada segundo es liviano."""
    cfg = _cfg_by_token(db, token)
    event = db.query(Event).filter(Event.id == cfg.event_id).first()
    last = db.query(RouletteDraw).filter_by(event_id=cfg.event_id).order_by(RouletteDraw.id.desc()).first()
    out = {"style": _style(cfg), "event_name": event.name, "latest_id": last.id if last else 0, "draw": None}
    if last and after is not None and last.id > after:
        names = [p["name"] for p in candidates(db, event)]
        _RNG.shuffle(names)
        out["draw"] = {"id": last.id, "label": last.label, "mode": last.mode, "winners": json.loads(last.winners_json), "pool": names[:POOL_SIZE]}
    return out


@public_router.get("/r/{token}/asset/{tenant_id}/{filename}")
async def display_asset(token: str, tenant_id: str, filename: str, db: Session = Depends(get_db)):
    cfg = _cfg_by_token(db, token)
    event = db.query(Event).filter(Event.id == cfg.event_id).first()
    safe = os.path.basename(filename)
    if tenant_id != event.tenant_id or os.path.splitext(safe)[1].lower() not in ALLOWED_IMAGE_EXT:
        raise HTTPException(status_code=404, detail="Archivo no encontrado")
    path = os.path.join("data", tenant_id, "badge_assets", safe)
    if not os.path.isfile(path):
        raise HTTPException(status_code=404, detail="Archivo no encontrado")
    return FileResponse(path)
