"""Documentos del Evento (ítem 8) y Legalizaciones / gastos (ítem 7) — reunión 2026-09-21.

Ambas son secciones a nivel de evento (coordinador+, `comercial` incluida: no son módulos operativos
de registro). Los archivos viven en `data/<tenant>/event_docs/<evento>/` y `.../event_expenses/<evento>/`
con un nombre interno aleatorio; el nombre original se conserva en la base para las descargas."""
import io
import os
import uuid
import zipfile
from datetime import datetime
from decimal import Decimal, InvalidOperation

import pandas as pd
from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse, StreamingResponse
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image
from sqlalchemy.orm import Session

from app.auth import get_event_for_staff, require_role
from app.database import get_db
from app.models import EventDocument, EventExpense, StaffUser

router = APIRouter()

DOC_EXT = {".pdf", ".doc", ".docx", ".xls", ".xlsx", ".csv", ".ppt", ".pptx", ".txt", ".rtf", ".odt", ".ods",
           ".png", ".jpg", ".jpeg", ".webp", ".gif"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".webp", ".gif"}
MAX_FILE_BYTES = 25_000_000


def _store(tenant_id: str, kind: str, event_id: int, upload_name: str, content: bytes) -> str:
    folder = os.path.join("data", tenant_id, kind, str(event_id))
    os.makedirs(folder, exist_ok=True)
    path = os.path.join(folder, f"{uuid.uuid4().hex}{os.path.splitext(upload_name)[1].lower()}")
    with open(path, "wb") as f:
        f.write(content)
    return path


def _remove(path: str) -> None:
    if path and os.path.isfile(path):
        os.remove(path)


def delete_event_files(db: Session, event_id: int) -> None:
    """Borra filas y archivos de documentos y gastos de un evento (lo usa delete_event)."""
    for doc in db.query(EventDocument).filter(EventDocument.event_id == event_id).all():
        _remove(doc.stored_path)
    for exp in db.query(EventExpense).filter(EventExpense.event_id == event_id).all():
        _remove(exp.evidence_path)
    db.query(EventDocument).filter(EventDocument.event_id == event_id).delete()
    db.query(EventExpense).filter(EventExpense.event_id == event_id).delete()


# ---------------------------------------------------------------- Documentos del Evento (ítem 8)

def _doc_json(d: EventDocument) -> dict:
    return {
        "id": d.id, "name": d.name, "description": d.description or "", "filename": d.original_filename,
        "size_bytes": d.size_bytes, "created_at": d.created_at.isoformat() if d.created_at else None,
    }


@router.get("/events/{event_id}/documents")
async def list_documents(event_id: int, db: Session = Depends(get_db), staff: StaffUser = Depends(require_role("coordinador"))):
    event = get_event_for_staff(event_id, db, staff)
    docs = db.query(EventDocument).filter(EventDocument.event_id == event.id).order_by(EventDocument.created_at.desc()).all()
    return [_doc_json(d) for d in docs]


@router.post("/events/{event_id}/documents")
async def add_document(
    event_id: int, name: str = Form(...), description: str = Form(""), file: UploadFile = File(...),
    db: Session = Depends(get_db), staff: StaffUser = Depends(require_role("coordinador")),
):
    event = get_event_for_staff(event_id, db, staff)
    if not name.strip():
        raise HTTPException(status_code=400, detail="El documento necesita un nombre o referencia")
    ext = os.path.splitext(file.filename or "")[1].lower()
    if ext not in DOC_EXT:
        raise HTTPException(status_code=400, detail=f"Formato no permitido — usa uno de: {', '.join(sorted(DOC_EXT))}")
    content = await file.read()
    if not content or len(content) > MAX_FILE_BYTES:
        raise HTTPException(status_code=400, detail="El archivo está vacío o pesa más de 25 MB")
    doc = EventDocument(
        event_id=event.id, name=name.strip(), description=description.strip() or None,
        original_filename=os.path.basename(file.filename), mime_type=file.content_type, size_bytes=len(content),
        stored_path=_store(event.tenant_id, "event_docs", event.id, file.filename, content),
        uploaded_by_id=staff.id, created_at=datetime.utcnow(),
    )
    db.add(doc)
    db.commit()
    return _doc_json(doc)


def _get_doc(db: Session, event_id: int, doc_id: int, staff: StaffUser) -> EventDocument:
    event = get_event_for_staff(event_id, db, staff)
    doc = db.query(EventDocument).filter(EventDocument.id == doc_id, EventDocument.event_id == event.id).first()
    if not doc or not os.path.isfile(doc.stored_path):
        raise HTTPException(status_code=404, detail="Documento no encontrado")
    return doc


@router.get("/events/{event_id}/documents/zip")
async def download_all_documents(event_id: int, db: Session = Depends(get_db), staff: StaffUser = Depends(require_role("coordinador"))):
    """Todos los documentos del evento en un ZIP (al finalizar el evento). Cada archivo se nombra
    "<referencia> - <archivo original>" y los nombres repetidos se numeran."""
    event = get_event_for_staff(event_id, db, staff)
    docs = db.query(EventDocument).filter(EventDocument.event_id == event.id).order_by(EventDocument.created_at).all()
    if not docs:
        raise HTTPException(status_code=404, detail="Este evento no tiene documentos")
    buffer, used = io.BytesIO(), set()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        for d in docs:
            if not os.path.isfile(d.stored_path):
                continue
            entry = f"{d.name} - {d.original_filename}".replace("/", "_").replace("\\", "_")
            base, n = entry, 1
            while entry in used:
                n += 1
                stem, ext = os.path.splitext(base)
                entry = f"{stem} ({n}){ext}"
            used.add(entry)
            zf.write(d.stored_path, entry)
    buffer.seek(0)
    return StreamingResponse(
        buffer, media_type="application/zip",
        headers={"Content-Disposition": f'attachment; filename="Documentos_{event.event_code}.zip"'},
    )


@router.get("/events/{event_id}/documents/{doc_id}/download")
async def download_document(event_id: int, doc_id: int, db: Session = Depends(get_db), staff: StaffUser = Depends(require_role("coordinador"))):
    doc = _get_doc(db, event_id, doc_id, staff)
    return FileResponse(doc.stored_path, filename=doc.original_filename, media_type=doc.mime_type or "application/octet-stream")


@router.delete("/events/{event_id}/documents/{doc_id}")
async def delete_document(event_id: int, doc_id: int, db: Session = Depends(get_db), staff: StaffUser = Depends(require_role("coordinador"))):
    doc = _get_doc(db, event_id, doc_id, staff)
    _remove(doc.stored_path)
    db.delete(doc)
    db.commit()
    return {"message": "Documento eliminado"}


# ---------------------------------------------------------------- Legalizaciones (ítem 7)

def _expense_json(e: EventExpense) -> dict:
    return {
        "id": e.id, "category": e.category, "responsible": e.responsible, "description": e.description or "",
        "applies_to": e.applies_to or "", "amount": float(e.amount or 0), "has_evidence": bool(e.evidence_path),
        "created_at": e.created_at.isoformat() if e.created_at else None,
    }


@router.get("/events/{event_id}/expenses")
async def list_expenses(event_id: int, db: Session = Depends(get_db), staff: StaffUser = Depends(require_role("coordinador"))):
    event = get_event_for_staff(event_id, db, staff)
    rows = db.query(EventExpense).filter(EventExpense.event_id == event.id).order_by(EventExpense.created_at).all()
    return {"expenses": [_expense_json(e) for e in rows], "total": float(sum((e.amount or 0) for e in rows))}


@router.post("/events/{event_id}/expenses")
async def add_expense(
    event_id: int, category: str = Form(...), responsible: str = Form(...), description: str = Form(""),
    applies_to: str = Form(""), amount: str = Form(...), evidence: UploadFile = File(None),
    db: Session = Depends(get_db), staff: StaffUser = Depends(require_role("coordinador")),
):
    event = get_event_for_staff(event_id, db, staff)
    if not category.strip() or not responsible.strip():
        raise HTTPException(status_code=400, detail="La categoría y la persona responsable son obligatorias")
    try:
        value = Decimal(amount.replace(",", ".").strip())
        if value < 0:
            raise InvalidOperation
    except InvalidOperation:
        raise HTTPException(status_code=400, detail="El valor del gasto debe ser un número mayor o igual a 0")

    evidence_path = evidence_name = None
    if evidence is not None and evidence.filename:
        ext = os.path.splitext(evidence.filename)[1].lower()
        if ext not in IMAGE_EXT:
            raise HTTPException(status_code=400, detail="La evidencia debe ser una imagen (PNG, JPG, WebP o GIF)")
        content = await evidence.read()
        if not content or len(content) > MAX_FILE_BYTES:
            raise HTTPException(status_code=400, detail="La imagen está vacía o pesa más de 25 MB")
        try:
            Image.open(io.BytesIO(content)).verify()
        except Exception:
            raise HTTPException(status_code=400, detail="La evidencia no es una imagen válida")
        evidence_path = _store(event.tenant_id, "event_expenses", event.id, evidence.filename, content)
        evidence_name = os.path.basename(evidence.filename)

    expense = EventExpense(
        event_id=event.id, category=category.strip(), responsible=responsible.strip(),
        description=description.strip() or None, applies_to=applies_to.strip() or None, amount=value,
        evidence_path=evidence_path, evidence_name=evidence_name, created_by_id=staff.id, created_at=datetime.utcnow(),
    )
    db.add(expense)
    db.commit()
    return _expense_json(expense)


def _get_expense(db: Session, event_id: int, expense_id: int, staff: StaffUser) -> EventExpense:
    event = get_event_for_staff(event_id, db, staff)
    exp = db.query(EventExpense).filter(EventExpense.id == expense_id, EventExpense.event_id == event.id).first()
    if not exp:
        raise HTTPException(status_code=404, detail="Gasto no encontrado")
    return exp


@router.get("/events/{event_id}/expenses/report")
async def expenses_report(event_id: int, db: Session = Depends(get_db), staff: StaffUser = Depends(require_role("coordinador"))):
    """Reporte Excel de legalizaciones: título del evento, una fila por gasto con su evidencia
    incrustada como imagen, y el total al final. Se puede bajar en cualquier momento (normalmente al
    finalizar el evento)."""
    event = get_event_for_staff(event_id, db, staff)
    rows = db.query(EventExpense).filter(EventExpense.event_id == event.id).order_by(EventExpense.created_at).all()
    headers = ["Fecha", "Categoría", "Responsable", "Descripción", "Aplica a", "Valor", "Evidencia"]
    df = pd.DataFrame(
        [[e.created_at.strftime("%Y-%m-%d") if e.created_at else "", e.category, e.responsible, e.description or "",
          e.applies_to or "", float(e.amount or 0), "" if e.evidence_path else "Sin evidencia"] for e in rows],
        columns=headers,
    )
    title_rows, header_row = 4, 5
    out = io.BytesIO()
    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Legalizaciones", startrow=header_row - 1)
        ws = writer.sheets["Legalizaciones"]
        for i, text in enumerate([
            f"Legalizaciones — {event.name} ({event.event_code})",
            f"Cliente/Cuenta: {event.tenant.name if event.tenant else event.tenant_id}",
            f"Total de gastos: {sum((e.amount or 0) for e in rows):,.2f}",
        ], start=1):
            ws.merge_cells(f"A{i}:G{i}")
            ws.cell(row=i, column=1, value=text).font = Font(bold=True, size=13 if i == 1 else 11, color="0A0E2E")
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=header_row, column=col)
            cell.fill = PatternFill(start_color="0A0E2E", end_color="0A0E2E", fill_type="solid")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center")
        for col, width in zip("ABCDEFG", (12, 20, 24, 44, 24, 16, 30)):
            ws.column_dimensions[col].width = width
        for offset, e in enumerate(rows, start=1):
            row = header_row + offset
            ws.cell(row=row, column=6).number_format = "#,##0.00"
            ws.cell(row=row, column=4).alignment = Alignment(wrap_text=True, vertical="top")
            if e.evidence_path and os.path.isfile(e.evidence_path):
                thumb = Image.open(e.evidence_path).convert("RGB")
                thumb.thumbnail((220, 165))
                buf = io.BytesIO()
                thumb.save(buf, "PNG")
                buf.seek(0)
                img = XLImage(buf)
                ws.add_image(img, f"G{row}")
                ws.row_dimensions[row].height = 130
        total_row = header_row + len(rows) + 1
        ws.cell(row=total_row, column=5, value="TOTAL").font = Font(bold=True)
        total_cell = ws.cell(row=total_row, column=6, value=float(sum((e.amount or 0) for e in rows)))
        total_cell.font, total_cell.number_format = Font(bold=True), "#,##0.00"
    out.seek(0)
    return StreamingResponse(
        out, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="Legalizaciones_{event.event_code}.xlsx"'},
    )


@router.get("/events/{event_id}/expenses/{expense_id}/evidence")
async def expense_evidence(event_id: int, expense_id: int, db: Session = Depends(get_db), staff: StaffUser = Depends(require_role("coordinador"))):
    exp = _get_expense(db, event_id, expense_id, staff)
    if not exp.evidence_path or not os.path.isfile(exp.evidence_path):
        raise HTTPException(status_code=404, detail="Este gasto no tiene evidencia")
    return FileResponse(exp.evidence_path, filename=exp.evidence_name)


@router.delete("/events/{event_id}/expenses/{expense_id}")
async def delete_expense(event_id: int, expense_id: int, db: Session = Depends(get_db), staff: StaffUser = Depends(require_role("coordinador"))):
    exp = _get_expense(db, event_id, expense_id, staff)
    _remove(exp.evidence_path)
    db.delete(exp)
    db.commit()
    return {"message": "Gasto eliminado"}
