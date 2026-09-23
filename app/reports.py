from app.timeutil import to_local
import os

import pandas as pd
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.models import User, AccessLog, EventAttendee, Event, EventFieldConfig

# Sprint 2.4 Fase 16 (2026-09-17, pedido explícito): reemplaza el reporte de formato fijo/legacy
# (columnas "Tipo_Pago"/"usuario"/etc, pensadas para un integrador externo que ya no aplica) por
# uno dinámico: solo las columnas que de verdad se usaron en ESTE evento, con el nombre real de
# cada opcional, más el "Tipo de Registro" (de dónde salió cada acreditación) y la hora exacta con
# segundos — pensado para mandarse directo al cliente, con formato de Excel real (encabezados con
# estilo, autofiltro, un título arriba con cliente/evento/código).
REGISTRATION_METHOD_LABELS = {
    "tradicional": "Tradicional",
    "autoregistro": "Autoregistro",
    "biometrico": "Biométrico",
    "qr": "QR",
}

# (clave del campo, encabezado). id/first_name/last_name son la identidad — siempre van, tengan
# dato o no (una fila sin nombre igual necesita esas 3 columnas para ubicarla). El resto solo
# aparece si AL MENOS una persona de este evento trae un valor real ahí.
BASE_FIELDS = [
    ("id", "Cédula"),
    ("first_name", "Nombres"),
    ("last_name", "Apellidos"),
    ("role", "Cargo"),
    ("entity", "Entidad"),
    ("phone", "Teléfono"),
    ("email", "Correo Electrónico"),
    ("opt_1", "Tipo de Asistente"),
    ("categories", "Categoría"),  # ítem 14: solo sale si alguien tiene categoría en este evento
    ("certificate", "Certificado"),  # ítem 5: solo sale si el evento tiene el módulo de certificados activo
]
ALWAYS_INCLUDED_BASE_KEYS = {"id", "first_name", "last_name"}
UPPERCASE_KEYS = {"first_name", "last_name", "entity", "role"}

HEADER_FILL = PatternFill(start_color="0A0E2E", end_color="0A0E2E", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(bold=True, size=13, color="0A0E2E")
ACCENT_FILL = PatternFill(start_color="D4AF37", end_color="D4AF37", fill_type="solid")


def _event_directory_users(db: Session, event_id: int, tenant_id: str):
    """Mismo universo que el Directorio/Estadísticas de ESE evento (EventAttendee ∪ AccessLog
    filtrados por event_id) — no todo el tenant. Réplica local de la misma lógica que
    app/routers/stats.py: _event_directory_users, para no cruzar imports entre app/reports.py y
    app/routers/* (capas separadas a propósito en este proyecto)."""
    attendee_ids = {a.user_id for a in db.query(EventAttendee).filter(EventAttendee.event_id == event_id)}
    log_ids = {l.user_id for l in db.query(AccessLog).filter(AccessLog.event_id == event_id)}
    all_ids = attendee_ids | log_ids
    if not all_ids:
        return []
    return db.query(User).filter(User.tenant_id == tenant_id, User.id.in_(all_ids)).all()


def _has_value(value) -> bool:
    if value is None:
        return False
    if isinstance(value, str):
        return value.strip() != ""
    return True


class ReportManager:
    @staticmethod
    def generate_excel_report(db: Session, event_id: int, tenant_id: str, output_path: str):
        event = db.query(Event).filter(Event.id == event_id).first()
        optional_labels = sorted(event.get_optional_labels().items(), key=lambda kv: int(kv[0].split("_")[1])) if event else []
        # Etiquetas propias de este evento (Parámetros, ítem 3a) — solo cambian el encabezado.
        config_rows = db.query(EventFieldConfig).filter(EventFieldConfig.event_id == event_id).all()
        custom = {r.field_key: r.label for r in config_rows if r.label}
        field_types = {r.field_key: r.field_type for r in config_rows}
        optional_labels = [(k, custom.get(k, label)) for k, label in optional_labels]
        users = _event_directory_users(db, event_id, tenant_id)

        # Primer log de REGISTRO real por persona (excluye "Actualizado" — esas son ediciones de
        # perfil, no una acreditación; si el primer log de alguien fuera un "Actualizado" viejo, el
        # reporte terminaría mostrando una hora de edición como si fuera la hora de llegada).
        logs = (
            db.query(AccessLog)
            .filter(AccessLog.event_id == event_id, AccessLog.record_type != "Actualizado")
            .order_by(AccessLog.user_id, AccessLog.timestamp)
            .all()
        )
        first_log_by_user = {}
        for log in logs:
            first_log_by_user.setdefault(log.user_id, log)

        categories_by_user = {a.user_id: ", ".join(a.get_categories()) for a in db.query(EventAttendee).filter(EventAttendee.event_id == event_id)}
        certificate_by_user = {a.user_id: bool(a.certificate) for a in db.query(EventAttendee).filter(EventAttendee.event_id == event_id)}
        certificates_on = bool(event and event.certificates_enabled)
        raw_rows = []
        for user in users:
            raw_rows.append({
                "id": user.id, "first_name": user.first_name, "last_name": user.last_name,
                "role": user.role, "entity": user.entity, "phone": user.phone,
                "email": user.email, "opt_1": user.opt_1, "categories": categories_by_user.get(user.id, ""),
                "certificate": ("Sí" if certificate_by_user.get(user.id) else "No") if certificates_on else "",
                "extras": user.get_extras(),
                "log": first_log_by_user.get(user.id),
            })

        # Bug real corregido (2026-09-17): un evento sin nadie en el directorio todavía (recién
        # creado, o con un roster que aún no llegó) da raw_rows=[] — ninguna columna "opcional"
        # calificaría como usada, así que solo quedan las de identidad. El bug de fondo era más
        # abajo: pd.DataFrame([]) (lista vacía de filas) no tiene columnas SIN IMPORTAR qué diga
        # esta lista, rompiendo get_column_letter(0) — hay que pasarle `columns=` explícito.
        base_columns = [
            (key, custom.get(key, label)) for key, label in BASE_FIELDS
            if key in ALWAYS_INCLUDED_BASE_KEYS or any(_has_value(row[key]) for row in raw_rows)
        ]
        def _signature_file(user_id, key):
            # Firmas (ítem 10): un archivo PNG por evento/persona/campo, ver routers/signatures.py.
            path = os.path.join("data", tenant_id, "signatures", str(event_id), f"{user_id}__{key}.png")
            return path if os.path.isfile(path) else None

        def _column_used(key):
            if field_types.get(key) == "signature":
                return any(_signature_file(row["id"], key) for row in raw_rows)
            return any(_has_value(row["extras"].get(key)) for row in raw_rows)

        optional_columns = [(key, label) for key, label in optional_labels if _column_used(key)]
        signature_keys = {key for key, _ in optional_columns if field_types.get(key) == "signature"}
        all_headers = [label for _, label in base_columns] + [label for _, label in optional_columns] + \
            ["Tipo de Registro", "Fecha de Registro", "Hora de Registro"]

        data = []
        for row in raw_rows:
            out = {}
            for key, label in base_columns:
                value = row[key] or ""
                out[label] = value.upper() if key in UPPERCASE_KEYS and isinstance(value, str) else value
            for key, label in optional_columns:
                if key in signature_keys:
                    out[label] = "Firmado" if _signature_file(row["id"], key) else ""
                elif field_types.get(key) == "boolean":
                    raw = str(row["extras"].get(key, "")).strip().lower()
                    out[label] = "Sí" if raw == "true" else "No" if raw == "false" else ""  # sin elegir = vacío
                elif field_types.get(key) == "consent":
                    out[label] = "Sí" if str(row["extras"].get(key, "")).strip().lower() == "true" else "No"
                else:
                    out[label] = row["extras"].get(key, "")

            log = row["log"]
            if log:
                out["Tipo de Registro"] = REGISTRATION_METHOD_LABELS.get(log.registration_method, log.registration_method or "Sin especificar")
                local = to_local(log.timestamp)  # la base guarda UTC; el reporte se lee en hora local (Sprint 4)
                out["Fecha de Registro"] = local.strftime("%Y-%m-%d")
                out["Hora de Registro"] = local.strftime("%H:%M:%S")
            else:
                out["Tipo de Registro"] = "No registrado"
                out["Fecha de Registro"] = ""
                out["Hora de Registro"] = ""
            data.append(out)

        df = pd.DataFrame(data, columns=all_headers)
        headers = list(df.columns)
        n_cols = len(headers)
        last_col_letter = get_column_letter(n_cols)

        TITLE_ROWS = 4  # cliente/cuenta, evento, código, fila en blanco
        header_row = TITLE_ROWS + 1

        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            df.to_excel(writer, index=False, sheet_name="Reporte", startrow=header_row - 1)
            ws = writer.sheets["Reporte"]

            # Título arriba de la tabla (2026-09-17, pedido explícito): nombre del cliente/cuenta,
            # nombre del evento y código — para que quien reciba el Excel sepa de qué es sin
            # tener que preguntar, ya que se manda directo al cliente.
            tenant_name = event.tenant.name if event and event.tenant else tenant_id
            title_lines = [
                f"Cliente/Cuenta: {tenant_name}",
                f"Evento: {event.name if event else ''}",
                f"Código: {event.event_code if event else ''}",
            ]
            for i, text in enumerate(title_lines, start=1):
                ws.merge_cells(f"A{i}:{last_col_letter}{i}")
                cell = ws.cell(row=i, column=1, value=text)
                cell.font = TITLE_FONT
                cell.alignment = Alignment(horizontal="left", vertical="center")
                if i == 1:
                    cell.fill = ACCENT_FILL

            # Encabezados con estilo + autofiltro + ancho de columna según el contenido real.
            for col_idx, header in enumerate(headers, start=1):
                cell = ws.cell(row=header_row, column=col_idx)
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT
                cell.alignment = Alignment(horizontal="center", vertical="center")

                max_len = len(str(header))
                for value in df[header]:
                    max_len = max(max_len, len(str(value)) if value is not None else 0)
                ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 3, 45)

            # Firmas: la imagen queda pegada en la celda de esa persona (fila más alta para verla).
            for key, label in optional_columns:
                if key not in signature_keys:
                    continue
                col_idx = headers.index(label) + 1
                ws.column_dimensions[get_column_letter(col_idx)].width = 24
                for offset, row in enumerate(raw_rows, start=1):
                    sig = _signature_file(row["id"], key)
                    if not sig:
                        continue
                    img = XLImage(sig)
                    img.width, img.height = 150, 56
                    ws.add_image(img, f"{get_column_letter(col_idx)}{header_row + offset}")
                    ws.row_dimensions[header_row + offset].height = 46

            last_data_row = header_row + len(df)
            ws.auto_filter.ref = f"A{header_row}:{last_col_letter}{last_data_row}"
            ws.freeze_panes = f"A{header_row + 1}"

        return output_path
