"""Campo «Pago» (Wompi) de los Formularios Web: precios con reglas y descuentos, firmas de Wompi, y el flujo completo
(la inscripción NO existe hasta que el pago se aprueba). Las llaves son inventadas para la prueba; Wompi real no se toca."""
import hashlib
import json
from datetime import date

import pytest

from app import formlib, wompi
from app.models import FormPayment
from tests.conftest import login
from tests.test_forms import _basic_fields, _create, _design, _event, _open, _put, _status, _submit, _url, _person_values  # noqa: F401

PROD = {"WOMPI_PUBLIC_KEY": "pub_prod_TEST", "WOMPI_INTEGRITY_SECRET": "prod_integrity_x", "WOMPI_EVENTS_SECRET": "prod_events_x"}
SANDBOX = {"WOMPI_SANDBOX_PUBLIC_KEY": "pub_test_TEST", "WOMPI_SANDBOX_INTEGRITY_SECRET": "test_integrity_x", "WOMPI_SANDBOX_EVENTS_SECRET": "test_events_x"}


@pytest.fixture(autouse=True)
def _no_dns(monkeypatch):
    monkeypatch.setattr("app.routers.forms_public.check_email", lambda e: (True, ""))


@pytest.fixture()
def keys(monkeypatch):
    for k, v in {**PROD, **SANDBOX}.items():
        monkeypatch.setenv(k, v)


def _pay(**pay):
    return {"id": "pago", "type": "payment", "label": "Pago de inscripción", "pay": {"mode": "fixed", "amount": 150000, **pay}}


TIPO = {"id": "tipo", "type": "select", "label": "Tipo de entrada", "options": ["General", "VIP", "Estudiante"]}


def _with_payment(client, ev, pay=None, extra=(), status="activo", **settings):
    form = _create(client, ev)
    fields = _basic_fields() + list(extra) + [pay or _pay()]
    r = _put(client, ev, form, design=_design(fields), **({"settings": settings} if settings else {}))
    assert r.status_code == 200, r.text
    _open(client, ev, form, status=status)
    return form


def _values(**more):
    return {**_person_values(), **more}


def _event_for(env, ref, cents, status, txn="T1", method="CARD"):
    ev = {"event": "transaction.updated", "environment": env, "timestamp": 1700000000, "sent_at": "2026-09-24T10:00:00Z",
          "data": {"transaction": {"id": txn, "status": status, "amount_in_cents": cents, "reference": ref, "currency": "COP", "payment_method_type": method}},
          "signature": {"properties": ["transaction.id", "transaction.status", "transaction.amount_in_cents"], "checksum": ""}}
    secret = PROD["WOMPI_EVENTS_SECRET"] if env == "prod" else SANDBOX["WOMPI_SANDBOX_EVENTS_SECRET"]
    ev["signature"]["checksum"] = hashlib.sha256(f"{txn}{status}{cents}1700000000{secret}".encode()).hexdigest()   # a mano, según la especificación de Wompi
    return ev


def _staff(client):
    login(client, "coord1")


# ------------------------------- precios (lógica pura) -------------------------------
def _sanitize(fields, pay):
    design = {"theme": {}, "fields": {f["id"]: f for f in fields + [{"id": "pago", "type": "payment", "label": "Pago", "pay": pay}]}, "rows": []}
    return formlib.sanitize_design(design, set())["fields"]["pago"]["pay"]


def test_payment_config_is_validated():
    base = [dict(TIPO)]
    assert _sanitize(base, {"mode": "fixed", "amount": 150000})["amount"] == 150000
    for bad in ({"mode": "fixed", "amount": 0}, {"mode": "fixed", "amount": 500}, {"mode": "fixed", "amount": 10_000_001}, {"mode": "fixed", "amount": "abc"}):
        with pytest.raises(ValueError):
            _sanitize(base, bad)
    with pytest.raises(ValueError):       # regla que apunta a un campo inexistente
        _sanitize(base, {"mode": "rules", "amount": 1500, "rules": [{"when": [{"field": "nope", "op": "equals", "value": "x"}], "amount": 2000}]})
    with pytest.raises(ValueError):       # descuento en porcentaje fuera de rango
        _sanitize(base, {"mode": "fixed", "amount": 150000, "discounts": [{"kind": "percent", "value": 150}]})
    with pytest.raises(ValueError):       # fecha inválida
        _sanitize(base, {"mode": "fixed", "amount": 150000, "discounts": [{"kind": "percent", "value": 10, "when": [{"field": "@date", "op": "before", "value": "mañana"}]}]})
    two = {"theme": {}, "rows": [], "fields": {"a": {"id": "a", "type": "payment", "pay": {"amount": 2000}}, "b": {"id": "b", "type": "payment", "pay": {"amount": 2000}}}}
    with pytest.raises(ValueError, match="un solo campo de pago"):
        formlib.sanitize_design(two, set())


def test_amount_rules_and_discounts_combine():
    fields = [dict(TIPO)]
    pay = _sanitize(fields, {"mode": "rules", "amount": 100000, "rules": [
        {"label": "VIP", "when": [{"field": "tipo", "op": "equals", "value": "VIP"}], "amount": 300000}],
        "discounts": [
            {"label": "Estudiante", "kind": "percent", "value": 20, "when": [{"field": "tipo", "op": "equals", "value": "Estudiante"}]},
            {"label": "Pronto pago", "kind": "percent", "value": 10, "when": [{"field": "@date", "op": "before", "value": "2026-10-01"}]},
            {"label": "Bono", "kind": "amount", "value": 5000, "when": [{"field": "tipo", "op": "equals", "value": "VIP"}]}]})
    early, late = date(2026, 9, 30), date(2026, 10, 1)
    assert formlib.compute_amount(pay, {"tipo": "General"}, early)["amount"] == 90000                 # 100000 -10% pronto pago
    assert formlib.compute_amount(pay, {"tipo": "General"}, late)["amount"] == 100000                # el descuento por fecha ya venció (la fecha límite no cuenta)
    assert formlib.compute_amount(pay, {"tipo": "VIP"}, late)["amount"] == 295000                    # regla VIP 300000 - bono 5000
    assert formlib.compute_amount(pay, {"tipo": "Estudiante"}, early)["amount"] == 72000             # 100000 -20% = 80000, -10% = 72000 (en orden, sobre el monto vigente)
    q = formlib.compute_amount(pay, {"tipo": "VIP"}, early)
    assert q["base"] == 300000 and [a["label"] for a in q["applied"]] == ["VIP", "Pronto pago", "Bono"] and q["amount"] == 265000


def test_discount_never_goes_below_zero_and_hidden_fields_cannot_change_the_price():
    fields = [dict(TIPO), {"id": "code", "type": "text_short", "label": "Código", "show_if": {"field": "tipo", "op": "equals", "value": "VIP"}}]
    pay = _sanitize(fields, {"mode": "fixed", "amount": 20000, "discounts": [
        {"label": "Cupón", "kind": "amount", "value": 50000, "when": [{"field": "code", "op": "equals", "value": "GRATIS"}]}]})
    design = {"theme": {}, "rows": [], "fields": {f["id"]: f for f in fields}}
    design["fields"] = formlib.sanitize_design({"theme": {}, "rows": [], "fields": {f["id"]: f for f in fields}}, set())["fields"]
    values = {"tipo": "General", "code": "GRATIS"}        # el campo «code» está oculto para «General»: no puede dar descuento
    assert formlib.compute_amount(pay, formlib.priced_values(design, values), date.today())["amount"] == 20000
    values["tipo"] = "VIP"
    assert formlib.compute_amount(pay, formlib.priced_values(design, values), date.today())["amount"] == 0


# ------------------------------- firmas de Wompi -------------------------------
def test_integrity_signature_and_event_checksum_follow_wompi_spec(monkeypatch, keys):
    assert wompi.integrity_signature("REF1", 15000000, "COP", "sec") == hashlib.sha256(b"REF115000000COPsec").hexdigest()
    ev = _event_for("prod", "REF1", 15000000, "APPROVED")
    assert wompi.verify_event(ev) is True
    ev["data"]["transaction"]["amount_in_cents"] = 100          # alterado: el checksum ya no cuadra
    assert wompi.verify_event(ev) is False
    monkeypatch.delenv("WOMPI_EVENTS_SECRET")
    monkeypatch.delenv("WOMPI_SANDBOX_EVENTS_SECRET")
    assert wompi.verify_event(ev) is None                        # sin secretos no se puede verificar


# ------------------------------- flujo completo -------------------------------
def test_submission_waits_for_payment_and_is_invisible_until_approved(client, factory, keys):
    ev = _event(client, factory, )
    form = _with_payment(client, ev, feed="realtime")
    r = _submit(client, ev, form, _values())
    assert r.status_code == 200, r.text
    body = r.json()
    pay = body["payment"]
    assert body["payment_required"] and pay["amount_in_cents"] == 15000000 and pay["currency"] == "COP" and pay["public_key"] == "pub_prod_TEST" and pay["test"] is False
    assert pay["reference"].startswith(f"GW-{ev.id}-{form['id']}-")
    assert pay["integrity"] == wompi.integrity_signature(pay["reference"], 15000000, "COP", PROD["WOMPI_INTEGRITY_SECRET"])
    assert "prod_integrity_x" not in json.dumps(body)            # el secreto nunca sale
    assert pay["customer"]["email"] == "ana@example.com" and pay["customer"]["fullName"] == "Ana Mora"
    # todavía NO hay inscripción: ni en la lista, ni en el conteo, ni en la analítica, ni cargada a la base del evento
    _staff(client)
    assert client.get(f"/api/events/{ev.id}/forms/{form['id']}/submissions").json()["rows"] == []
    assert client.get(f"/api/events/{ev.id}/forms").json()[0]["submissions"] == 0
    assert client.get(f"/api/events/{ev.id}/forms/{form['id']}/analytics").json()["kpis"][0]["value"] == 0
    assert all(u["id"] != "1001" for u in client.get(f"/api/users?event_id={ev.id}").json())
    client.post("/logout")
    # Wompi confirma por webhook -> ahora sí
    event = _event_for("prod", pay["reference"], 15000000, "APPROVED", txn="TX-9")
    assert client.post("/webhooks/wompi", json=event).json() == {"ok": True}
    _staff(client)
    rows = client.get(f"/api/events/{ev.id}/forms/{form['id']}/submissions").json()["rows"]
    assert len(rows) == 1 and rows[0]["paid"] == 150000
    assert any(u["id"] == "1001" for u in client.get(f"/api/users?event_id={ev.id}").json())      # carga en tiempo real
    payments = client.get(f"/api/events/{ev.id}/forms/{form['id']}/payments").json()
    assert payments["total_cop"] == 150000 and payments["rows"][0]["status"] == "approved" and payments["rows"][0]["transaction_id"] == "TX-9" and not payments["rows"][0]["orphan"]
    client.post("/logout")
    # repetir el mismo webhook no duplica nada
    client.post("/webhooks/wompi", json=event)
    _staff(client)
    assert len(client.get(f"/api/events/{ev.id}/forms/{form['id']}/submissions").json()["rows"]) == 1


def test_declined_payment_leaves_no_registration_and_can_still_be_approved_later(client, factory, keys):
    ev = _event(client, factory)
    form = _with_payment(client, ev)
    ref = _submit(client, ev, form, _values()).json()["payment"]["reference"]
    assert client.post("/webhooks/wompi", json=_event_for("prod", ref, 15000000, "DECLINED")).status_code == 200
    _staff(client)
    assert client.get(f"/api/events/{ev.id}/forms/{form['id']}/submissions").json()["rows"] == []
    assert client.get(f"/api/events/{ev.id}/forms/{form['id']}/payments").json()["rows"][0]["status"] == "declined"
    client.post("/logout")
    # Wompi deja reintentar con la misma referencia: una aprobación posterior sí confirma
    client.post("/webhooks/wompi", json=_event_for("prod", ref, 15000000, "APPROVED", txn="T2"))
    _staff(client)
    assert len(client.get(f"/api/events/{ev.id}/forms/{form['id']}/submissions").json()["rows"]) == 1


def test_webhook_rejects_bad_signature_wrong_amount_and_wrong_environment(client, factory, keys):
    ev = _event(client, factory)
    form = _with_payment(client, ev)
    ref = _submit(client, ev, form, _values()).json()["payment"]["reference"]
    bad = _event_for("prod", ref, 15000000, "APPROVED")
    bad["signature"]["checksum"] = "0" * 64
    assert client.post("/webhooks/wompi", json=bad).status_code == 401
    wrong_amount = _event_for("prod", ref, 100, "APPROVED")
    assert client.post("/webhooks/wompi", json=wrong_amount).json()["ignored"] == "monto distinto al esperado"
    sandbox_event = _event_for("test", ref, 15000000, "APPROVED")             # firmado bien, pero de pruebas: no puede aprobar un pago real
    assert client.post("/webhooks/wompi", json=sandbox_event).json()["ignored"] == "ambiente distinto al del pago"
    assert client.post("/webhooks/wompi", json=_event_for("prod", "GW-otro", 100, "APPROVED")).json()["ignored"] == "referencia desconocida"
    _staff(client)
    assert client.get(f"/api/events/{ev.id}/forms/{form['id']}/submissions").json()["rows"] == []


def test_webhook_needs_configured_secrets(client, monkeypatch):
    for k in list(PROD) + list(SANDBOX):
        monkeypatch.delenv(k, raising=False)
    assert client.post("/webhooks/wompi", json={"event": "transaction.updated"}).status_code == 503


def test_browser_confirmation_is_checked_against_wompi_not_trusted(client, factory, keys, monkeypatch):
    ev = _event(client, factory)
    form = _with_payment(client, ev)
    body = _submit(client, ev, form, _values()).json()
    ref, token = body["payment"]["reference"], body["pay_token"]
    calls = []

    def fake_fetch(cfg, txn_id):
        calls.append(txn_id)
        return {"id": txn_id, "status": "APPROVED", "reference": "OTRA-REFERENCIA", "amount_in_cents": 15000000, "currency": "COP"} if txn_id == "forged" else \
               {"id": txn_id, "status": "APPROVED", "reference": ref, "amount_in_cents": 15000000, "currency": "COP", "payment_method_type": "PSE"}

    monkeypatch.setattr("app.routers.form_payments.wompi.fetch_transaction", fake_fetch)
    url = f"{_url(ev, form)}/pay/confirm"
    assert client.post(url, json={"pt": token, "transaction_id": "forged"}).json()["status"] == "pending"     # otra referencia: no cuenta
    assert client.post(url, json={"pt": "token-falso", "transaction_id": "x"}).status_code == 403
    r = client.post(url, json={"pt": token, "transaction_id": "good"}).json()
    assert r["status"] == "approved" and r["thanks"]["title"]
    assert client.get(f"{_url(ev, form)}/pay/status", params={"pt": token}).json()["status"] == "approved"
    _staff(client)
    assert len(client.get(f"/api/events/{ev.id}/forms/{form['id']}/submissions").json()["rows"]) == 1


def test_no_payment_config_means_no_half_saved_registration(client, factory, monkeypatch):
    for k in list(PROD) + list(SANDBOX):
        monkeypatch.delenv(k, raising=False)
    ev = _event(client, factory)
    form = _with_payment(client, ev)
    r = _submit(client, ev, form, _values())
    assert r.status_code == 503
    _staff(client)
    assert client.get(f"/api/events/{ev.id}/forms/{form['id']}/analytics").json()["kpis"][0]["value"] == 0


def test_price_from_rules_and_free_when_discount_covers_everything(client, factory, keys):
    ev = _event(client, factory)
    pay = _pay(mode="rules", amount=100000, rules=[{"label": "VIP", "when": [{"field": "tipo", "op": "equals", "value": "VIP"}], "amount": 300000}],
               discounts=[{"label": "Invitado", "kind": "percent", "value": 100, "when": [{"field": "tipo", "op": "equals", "value": "Estudiante"}]}])
    form = _with_payment(client, ev, pay=pay, extra=[TIPO])
    q = client.post(f"{_url(ev, form)}/quote", json={"values": {"tipo": "VIP"}}).json()
    assert q == {"has_payment": True, "amount": 300000, "base": 300000, "applied": [{"label": "VIP", "kind": "rule", "effect": "monto $300.000"}], "description": ""}
    assert client.post(f"{_url(ev, form)}/quote", json={"values": {"tipo": "General"}}).json()["amount"] == 100000
    # el servidor calcula el precio: lo que el navegador diga no cambia nada
    r = _submit(client, ev, form, _values(tipo="VIP"), sid="a", amount=1).json()
    assert r["payment"]["amount_in_cents"] == 30000000
    # cubierto al 100 %: no hay nada que cobrar y la inscripción se confirma de una vez
    free = _submit(client, ev, form, {"cedula": "2002", "nombres": "Luis", "apellidos": "Paz", "correo": "luis@example.com", "tipo": "Estudiante"}, sid="b")
    assert free.status_code == 200 and "payment_required" not in free.json() and free.json()["thanks"]
    _staff(client)
    assert len(client.get(f"/api/events/{ev.id}/forms/{form['id']}/submissions").json()["rows"]) == 1


def test_capacity_counts_people_who_are_paying_and_releases_it_when_they_fail(client, factory, keys):
    ev = _event(client, factory)
    form = _with_payment(client, ev)
    _staff(client)
    _status(client, ev, form, capacity=1)
    client.post("/logout")
    first = _submit(client, ev, form, _values(), sid="a").json()
    other = {"cedula": "3003", "nombres": "Eva", "apellidos": "Ruiz", "correo": "eva@example.com"}
    assert _submit(client, ev, form, other, sid="b").status_code == 409                    # el cupo lo tiene quien está pagando
    client.post("/webhooks/wompi", json=_event_for("prod", first["payment"]["reference"], 15000000, "DECLINED"))
    assert _submit(client, ev, form, other, sid="b").status_code == 200                    # rechazado: el cupo se libera


def test_retrying_replaces_the_previous_attempt_of_the_same_person(client, factory, keys):
    ev = _event(client, factory)
    form = _with_payment(client, ev)
    a = _submit(client, ev, form, _values(), sid="s1").json()["payment"]["reference"]
    b = _submit(client, ev, form, _values(), sid="s1").json()["payment"]["reference"]
    assert a != b
    client.post("/webhooks/wompi", json=_event_for("prod", b, 15000000, "APPROVED"))
    _staff(client)
    assert len(client.get(f"/api/events/{ev.id}/forms/{form['id']}/submissions").json()["rows"]) == 1


def test_test_link_uses_sandbox_keys_and_test_payments_stay_out_of_official_numbers(client, factory, keys):
    ev = _event(client, factory)
    form = _with_payment(client, ev, status="pruebas")
    _staff(client)
    key = client.get(f"/api/events/{ev.id}/forms/{form['id']}").json()["test_key"]
    client.post("/logout")
    body = _submit(client, ev, form, _values(), key=key).json()
    assert body["payment"]["test"] is True and body["payment"]["public_key"] == "pub_test_TEST"
    client.post("/webhooks/wompi", json=_event_for("test", body["payment"]["reference"], 15000000, "APPROVED"))
    _staff(client)
    base = f"/api/events/{ev.id}/forms/{form['id']}"
    assert client.get(f"{base}/submissions", params={"include_tests": False}).json()["rows"] == []
    assert client.get(f"{base}/payments").json()["rows"] == []                                  # el pago de pruebas no cuenta como ingreso real
    assert len(client.get(f"{base}/payments", params={"include_tests": True}).json()["rows"]) == 1


def test_form_with_approved_payments_cannot_be_deleted_and_reports_show_the_money(client, factory, keys):
    from openpyxl import load_workbook
    import io
    ev = _event(client, factory)
    pay = _pay(discounts=[{"label": "Pronto pago", "kind": "percent", "value": 10, "when": [{"field": "@date", "op": "on_or_after", "value": "2020-01-01"}]}])
    form = _with_payment(client, ev, pay=pay)
    ref = _submit(client, ev, form, _values()).json()["payment"]["reference"]
    client.post("/webhooks/wompi", json=_event_for("prod", ref, 13500000, "APPROVED", txn="TX-1", method="PSE"))
    _staff(client)
    base = f"/api/events/{ev.id}/forms/{form['id']}"
    r = client.delete(base)
    assert r.status_code == 409 and "pagos aprobados" in r.json()["detail"]
    ws = load_workbook(io.BytesIO(client.get(f"{base}/report").content)).active
    header = [c.value for c in ws[2]]
    row = [c.value for c in ws[3]]
    assert "Monto pagado (COP)" in header and row[header.index("Monto pagado (COP)")] == 135000
    assert row[header.index("Referencia de pago")] == ref and row[header.index("Método de pago")] == "PSE"
    assert "Pronto pago" in row[header.index("Reglas y descuentos")]
    kpis = {k["label"]: k["value"] for k in client.get(f"{base}/analytics").json()["kpis"]}
    assert kpis["Ingresos (aprobados)"] == "$135.000"


def test_payment_field_state_is_reported_to_the_editor(client, factory, keys, monkeypatch):
    ev = _event(client, factory)
    form = _with_payment(client, ev)
    _staff(client)
    assert client.get(f"/api/events/{ev.id}/forms/{form['id']}").json()["payments"] == {"has_field": True, "sandbox_configured": True, "production_configured": True}
    for k in SANDBOX:
        monkeypatch.delenv(k)
    assert client.get(f"/api/events/{ev.id}/forms/{form['id']}").json()["payments"]["sandbox_configured"] is False


# ------------------------------- el ambiente lo dice la llave, no el nombre de la variable -------------------------------
def _only_main(monkeypatch, public):
    for k in list(PROD) + list(SANDBOX):
        monkeypatch.delenv(k, raising=False)
    monkeypatch.setenv("WOMPI_PUBLIC_KEY", public)
    monkeypatch.setenv("WOMPI_INTEGRITY_SECRET", "integ")
    monkeypatch.setenv("WOMPI_EVENTS_SECRET", "evt")


def test_test_keys_under_the_main_variables_are_treated_as_sandbox(client, factory, monkeypatch):
    _only_main(monkeypatch, "pub_test_MAIN")
    assert wompi.config(False)["test"] is True and wompi.config(False)["api"].startswith("https://sandbox.")
    ev = _event(client, factory)
    form = _with_payment(client, ev)                                   # formulario ACTIVO, pero con llaves de pruebas
    body = _submit(client, ev, form, _values()).json()
    assert body["payment"]["test"] is True and body["payment"]["public_key"] == "pub_test_MAIN"
    _staff(client)
    assert client.get(f"/api/events/{ev.id}/forms/{form['id']}").json()["payments"]["production_configured"] is False   # no cuenta como producción
    assert client.get(f"/api/events/{ev.id}/forms/{form['id']}/payments").json()["rows"] == []                          # y no suma como ingreso real


def test_a_form_in_pruebas_never_uses_real_keys(client, factory, monkeypatch):
    _only_main(monkeypatch, "pub_prod_REAL")
    assert wompi.config(True) is None                                   # solo hay llaves reales: el modo pruebas no puede cobrar
    ev = _event(client, factory)
    form = _with_payment(client, ev, status="pruebas")
    _staff(client)
    key = client.get(f"/api/events/{ev.id}/forms/{form['id']}").json()["test_key"]
    client.post("/logout")
    assert _submit(client, ev, form, _values(), key=key).status_code == 503


# ------------------------------- precio de referencia en USD/EUR y traducción -------------------------------
def test_reference_rates_endpoint_and_fx_fallback(client, factory, keys, monkeypatch):
    from app import fx
    ev = _event(client, factory)
    form = _with_payment(client, ev)
    good = {"base": "COP", "rates": {"USD": 0.0003, "EUR": 0.00027}, "as_of": "Thu, 24 Sep 2026", "source": "exchangerate-api.com"}
    monkeypatch.setattr(fx, "_fetch", lambda: good)
    monkeypatch.setitem(fx._cache, "data", None)
    monkeypatch.setitem(fx._cache, "at", 0.0)
    assert client.get(f"{_url(ev, form)}/rates").json()["rates"] == good["rates"]
    monkeypatch.setattr(fx, "_fetch", lambda: None)                       # la fuente cae: se sigue con la última tasa buena (< 48 h)
    monkeypatch.setitem(fx._cache, "at", fx._cache["at"] - 7 * 3600)
    assert client.get(f"{_url(ev, form)}/rates").json()["rates"] == good["rates"]
    monkeypatch.setitem(fx._cache, "at", fx._cache["at"] - 50 * 3600)     # demasiado vieja: se oculta el selector, el cobro sigue en COP
    assert client.get(f"{_url(ev, form)}/rates").json()["rates"] == {}


def test_fx_ignores_garbage_from_the_source(monkeypatch):
    from app import fx
    class R:
        def __init__(self, body): self.body = body
        def read(self): return self.body
        def __enter__(self): return self
        def __exit__(self, *a): return False
    monkeypatch.setattr(fx.urllib.request, "urlopen", lambda *a, **k: R(b'{"result":"success","rates":{"USD":0,"EUR":1}}'))
    assert fx._fetch() is None                                            # una tasa en 0 nunca debe usarse
    monkeypatch.setattr(fx.urllib.request, "urlopen", lambda *a, **k: R(b'{"result":"success","rates":{"USD":0.0003,"EUR":0.00027},"time_last_update_utc":"Thu, 24 Sep 2026 00:02:32 +0000"}'))
    assert fx._fetch()["rates"] == {"USD": 0.0003, "EUR": 0.00027}


def test_language_and_translate_settings_reach_the_public_page(client, factory, keys):
    ev = _event(client, factory)
    form = _with_payment(client, ev)
    assert client.get(f"{_url(ev, form)}/state").json()["ui"] == {"language": "es", "translate": True}
    _staff(client)
    assert _put(client, ev, form, settings={"language": "en", "translate": False}).json()["settings"]["language"] == "en"
    assert _put(client, ev, form, settings={"language": "klingon"}).json()["settings"]["language"] == "es"      # idioma inválido: vuelve al predeterminado
    _put(client, ev, form, settings={"language": "pt", "translate": False})
    client.post("/logout")
    assert client.get(f"{_url(ev, form)}/state").json()["ui"] == {"language": "pt", "translate": False}
